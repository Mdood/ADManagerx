from __future__ import annotations

from dataclasses import dataclass
from typing import Optional, Iterable, Sequence, List, Any
import ssl
import json
import time
import threading
from datetime import datetime, timezone, timedelta

import winrm
from openpyxl import load_workbook
from ldap3 import (
    ALL,
    ALL_ATTRIBUTES,
    BASE,
    LEVEL,
    SUBTREE,
    MODIFY_ADD,
    MODIFY_DELETE,
    MODIFY_REPLACE,
    Connection,
    Server,
    Tls,
)

from .models import LdapSettings


_SEARCH_CACHE = {}
_CACHE_TTL = 45
_ACTIVE_SETTINGS = threading.local()

session = None


def set_current_ldap_settings(settings_id):
    """Set the LDAP settings row used by this request/thread."""
    try:
        _ACTIVE_SETTINGS.settings_id = int(settings_id) if settings_id else None
    except Exception:
        _ACTIVE_SETTINGS.settings_id = None


def clear_current_ldap_settings():
    _ACTIVE_SETTINGS.settings_id = None


def get_current_ldap_settings_id():
    return getattr(_ACTIVE_SETTINGS, "settings_id", None)


class ADServiceError(Exception):
    pass


@dataclass
class ADConfig:
    server_uri: str
    use_ssl: bool
    bind_dn: str
    bind_password: str
    base_dn: str
    users_ou_dn: str
    computers_ou_dn: str
    groups_ou_dn: str
    upn_suffix: str
    user_search_filter: str
    safe_mode: bool
    server_name: str


# -----------------------------
# Generic safety helpers
# -----------------------------

def _safe_str(value: Any) -> str:
    if value is None:
        return ""
    try:
        return str(value).strip()
    except Exception:
        return ""


def _safe_lower(value: Any) -> str:
    return _safe_str(value).lower()


def _entry_value(entry, *attr_names: str, default=None):
    for attr_name in attr_names:
        try:
            if attr_name in entry:
                value = getattr(entry, attr_name).value
                if value is not None:
                    return value
        except Exception:
            continue
    return default


def _entry_str(entry, *attr_names: str, default="") -> str:
    return _safe_str(_entry_value(entry, *attr_names, default=default))


def _entry_list(entry, attr_name: str) -> list[str]:
    try:
        if attr_name not in entry:
            return []
        values = getattr(entry, attr_name).values
        return [_safe_str(v) for v in values if _safe_str(v)]
    except Exception:
        return []


def _cache_get(key):
    item = _SEARCH_CACHE.get(key)
    if not item:
        return None

    ts, value = item
    if time.time() - ts > _CACHE_TTL:
        _SEARCH_CACHE.pop(key, None)
        return None

    return value


def _cache_set(key, value):
    _SEARCH_CACHE[key] = (time.time(), value)


def _escape_ldap_filter_value(value: str) -> str:
    value = str(value or "")
    return (
        value.replace("\\", "\\5c")
        .replace("*", "\\2a")
        .replace("(", "\\28")
        .replace(")", "\\29")
        .replace("\x00", "\\00")
    )


def _ps_escape_single_quotes(val) -> str:
    return str(val or "").replace("'", "''")


def _username_to_upn(username: str, upn_suffix: str) -> str:
    if "@" in username:
        return username
    suffix = _safe_str(upn_suffix)
    return f"{username}@{suffix}" if suffix else username


def _normalize_computer_sam(name: str) -> str:
    name = _safe_str(name)
    return name if name.endswith("$") else f"{name}$"


def _normalize_group_scope(value: str) -> str:
    value = _safe_lower(value)
    if not value:
        return ""

    mapping = {
        "global": "Global",
        "domainlocal": "DomainLocal",
        "domain local": "DomainLocal",
        "local": "DomainLocal",
        "universal": "Universal",
    }
    if value not in mapping:
        raise ADServiceError(f"Invalid group scope: {value}")
    return mapping[value]


def _normalize_group_category(value: str) -> str:
    value = _safe_lower(value)
    if not value:
        return ""

    mapping = {
        "security": "Security",
        "distribution": "Distribution",
    }
    if value not in mapping:
        raise ADServiceError(f"Invalid group category: {value}")
    return mapping[value]


def _filetime_to_int(value) -> int:
    try:
        if value in (None, "", 0, "0"):
            return 0
        return int(value)
    except Exception:
        return 0


def _windows_filetime_to_datetime(value) -> Optional[datetime]:
    raw = _filetime_to_int(value)
    if raw <= 0:
        return None

    if raw in (9223372036854775807, 9223372036854775808):
        return None

    try:
        return datetime.fromtimestamp((raw - 116444736000000000) / 10000000, tz=timezone.utc)
    except Exception:
        return None


def _datetime_to_display(value: Optional[datetime]) -> str:
    if not value:
        return ""
    try:
        return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    except Exception:
        return ""


def _now_utc() -> datetime:
    return datetime.now(timezone.utc)


def _manager_display_from_dn(manager_dn: str) -> str:
    manager_dn = _safe_str(manager_dn)
    if not manager_dn:
        return ""

    first_rdn = manager_dn.split(",", 1)[0].strip()
    if "=" in first_rdn:
        return first_rdn.split("=", 1)[1].strip()
    return manager_dn


def _best_name_from_entry(entry, default="") -> str:
    return (
        _entry_str(entry, "displayName")
        or _entry_str(entry, "cn")
        or _entry_str(entry, "name")
        or _entry_str(entry, "ou")
        or _entry_str(entry, "sAMAccountName")
        or default
    )


def _dedupe_dicts_by_key(items: list[dict], key_name: str) -> list[dict]:
    seen = set()
    out = []
    for item in items:
        key = _safe_str(item.get(key_name))
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(item)
    return out


def _chunked(values: Sequence[str], size: int = 100) -> Iterable[List[str]]:
    items = list(values)
    for i in range(0, len(items), size):
        yield items[i:i + size]


# -----------------------------
# Config / connection
# -----------------------------

def _load_config(settings_id=None) -> ADConfig:
    global session

    active_settings_id = settings_id if settings_id is not None else get_current_ldap_settings_id()
    cfg_db = LdapSettings.get_settings(settings_id=active_settings_id)
    if not cfg_db or not LdapSettings.is_configured(settings_id=getattr(cfg_db, "id", None)):
        raise ADServiceError("LDAP settings are not configured.")

    bind_user = _safe_str(cfg_db.bind_dn)
    upn_suffix = _safe_str(cfg_db.upn_suffix)

    if bind_user and "@" not in bind_user and "," not in bind_user and upn_suffix:
        bind_user = f"{bind_user}@{upn_suffix}"

    server_name = _safe_str(cfg_db.server_name)

    if server_name and bind_user and cfg_db.bind_password:
        session = winrm.Session(
            server_name,
            auth=(bind_user, cfg_db.bind_password),
            transport="ntlm",
        )
    else:
        session = None

    return ADConfig(
        server_uri=_safe_str(cfg_db.server_uri),
        use_ssl=bool(cfg_db.use_ssl),
        bind_dn=_safe_str(cfg_db.bind_dn),
        bind_password=_safe_str(cfg_db.bind_password),
        base_dn=_safe_str(cfg_db.base_dn),
        users_ou_dn=_safe_str(cfg_db.users_ou_dn),
        computers_ou_dn=_safe_str(cfg_db.computers_ou_dn),
        groups_ou_dn=_safe_str(cfg_db.groups_ou_dn),
        upn_suffix=upn_suffix,
        user_search_filter=_safe_str(cfg_db.user_search_filter) or "(sAMAccountName={username})",
        safe_mode=bool(cfg_db.safe_mode),
        server_name=server_name,
    )


def _connect(cfg: ADConfig) -> Connection:
    tls = Tls(validate=ssl.CERT_NONE) if cfg.use_ssl else None
    server = Server(cfg.server_uri, use_ssl=cfg.use_ssl, get_info=ALL, tls=tls)

    bind_user = _safe_str(cfg.bind_dn)
    if bind_user and "@" not in bind_user and "," not in bind_user and cfg.upn_suffix:
        bind_user = f"{bind_user}@{cfg.upn_suffix}"

    return Connection(server, user=bind_user, password=cfg.bind_password, auto_bind=True)


def _search_one(
    conn: Connection,
    base_dn: str,
    search_filter: str,
    attributes: Optional[List[str]] = None,
):
    attrs = attributes or ALL_ATTRIBUTES
    ok = conn.search(
        base_dn,
        search_filter,
        search_scope=SUBTREE,
        attributes=attrs,
        size_limit=1,
    )
    if not ok or not conn.entries:
        return None
    return conn.entries[0]


def _resolve_base_dn(conn: Connection, cfg: ADConfig) -> str:
    if _safe_str(cfg.base_dn):
        return _safe_str(cfg.base_dn)

    try:
        conn.search(
            search_base="",
            search_filter="(objectClass=*)",
            search_scope=BASE,
            attributes=["defaultNamingContext"],
            size_limit=1,
        )
        if conn.entries:
            value = _entry_str(conn.entries[0], "defaultNamingContext")
            if value:
                return value
    except Exception:
        pass

    raise ADServiceError("Cannot resolve base DN. Set base_dn in LDAP settings or ensure RootDSE is readable.")


def _run_ps(ps_script: str) -> str:
    global session

    if session is None:
        raise ADServiceError("WinRM session not initialized. Check LDAP/WinRM settings.")

    r = session.run_ps(ps_script)

    out = (r.std_out or b"").decode(errors="ignore")
    err = (r.std_err or b"").decode(errors="ignore")

    if r.status_code != 0:
        raise ADServiceError(
            f"WinRM PowerShell failed (status={r.status_code}). STDERR={err} STDOUT={out}"
        )

    return out


def _parse_json_output(output: str):
    if not output:
        return []

    output = output.strip()
    if not output:
        return []

    try:
        data = json.loads(output)
        if isinstance(data, dict):
            return [data]
        if isinstance(data, list):
            return data
        return []
    except json.JSONDecodeError as e:
        raise ADServiceError(f"Failed to parse PowerShell JSON output: {output}") from e


def test_connection(settings_id=None) -> None:
    cfg = _load_config(settings_id=settings_id)
    conn = _connect(cfg)
    conn.unbind()



def list_ldap_settings() -> list[dict]:
    return [
        {
            "id": obj.id,
            "name": obj.name or obj.domain_name or obj.user_domain or obj.server_uri,
            "domain_name": obj.domain_name or obj.user_domain,
            "server_uri": obj.server_uri,
            "is_default": obj.is_default,
            "is_active": obj.is_active,
        }
        for obj in LdapSettings.available_settings()
    ]

def test_ldap_health(settings_id=None) -> dict:
    """
    Lightweight LDAP health check.

    This only opens an LDAP connection, resolves the Base DN,
    then closes the connection. It does not modify Active Directory.
    """
    cfg = _load_config(settings_id=settings_id)

    result = {
        "status": "Unknown",
        "server_uri": cfg.server_uri,
        "base_dn": cfg.base_dn,
        "response_ms": 0,
        "error": "",
    }

    start = time.time()

    try:
        conn = _connect(cfg)
        try:
            resolved_base = _resolve_base_dn(conn, cfg)
            result["base_dn"] = resolved_base
            result["status"] = "Connected"
        finally:
            conn.unbind()

    except Exception as exc:
        result["status"] = "Failed"
        result["error"] = str(exc)

    finally:
        result["response_ms"] = int((time.time() - start) * 1000)

    return result


def test_winrm_health(settings_id=None) -> dict:
    """
    WinRM health check.

    This runs a tiny read-only PowerShell command through WinRM.
    It is heavier than LDAP because it uses remote PowerShell.
    """
    cfg = _load_config(settings_id=settings_id)

    result = {
        "status": "Unknown",
        "server_name": cfg.server_name,
        "safe_mode": cfg.safe_mode,
        "response_ms": 0,
        "error": "",
    }

    if cfg.safe_mode:
        result["status"] = "Safe Mode"
        result["error"] = "Safe mode is enabled, so WinRM commands are skipped."
        return result

    if not cfg.server_name:
        result["status"] = "Not Configured"
        result["error"] = "WinRM server_name is not configured."
        return result

    if session is None:
        result["status"] = "Not Initialized"
        result["error"] = "WinRM session is not initialized. Check LDAP/WinRM settings."
        return result

    start = time.time()

    try:
        output = _run_ps("Write-Output 'WINRM_OK=1'")

        if "WINRM_OK=1" in output:
            result["status"] = "Connected"
        else:
            result["status"] = "Unknown"
            result["error"] = "WinRM command completed but did not return the expected response."

    except Exception as exc:
        result["status"] = "Failed"
        result["error"] = str(exc)

    finally:
        result["response_ms"] = int((time.time() - start) * 1000)

    return result



# -----------------------------
# LDAP read helpers
# -----------------------------

def _get_user_dn(conn: Connection, cfg: ADConfig, username: str) -> Optional[str]:
    username = _safe_str(username)
    if not username:
        return None

    base_dn = _resolve_base_dn(conn, cfg)

    try:
        flt = cfg.user_search_filter.format(username=_escape_ldap_filter_value(username))
    except Exception:
        flt = f"(sAMAccountName={_escape_ldap_filter_value(username)})"

    entry = _search_one(conn, base_dn, flt, attributes=["distinguishedName"])
    if entry:
        return entry.entry_dn

    fallback_filter = (
        f"(&"
        f"(objectClass=user)"
        f"(!(objectClass=computer))"
        f"(|"
        f"(sAMAccountName={_escape_ldap_filter_value(username)})"
        f"(userPrincipalName={_escape_ldap_filter_value(username)})"
        f"(cn={_escape_ldap_filter_value(username)})"
        f")"
        f")"
    )
    entry = _search_one(conn, base_dn, fallback_filter, attributes=["distinguishedName"])
    if not entry:
        return None
    return entry.entry_dn


def _get_computer_dn(conn: Connection, cfg: ADConfig, computer_name: str) -> Optional[str]:
    computer_name = _safe_str(computer_name)
    if not computer_name:
        return None

    sam = _normalize_computer_sam(computer_name)
    base_dn = _resolve_base_dn(conn, cfg)

    search_filter = (
        f"(&"
        f"(objectClass=computer)"
        f"(|"
        f"(sAMAccountName={_escape_ldap_filter_value(sam)})"
        f"(cn={_escape_ldap_filter_value(computer_name)})"
        f"(name={_escape_ldap_filter_value(computer_name)})"
        f")"
        f")"
    )
    entry = _search_one(conn, base_dn, search_filter, attributes=["distinguishedName"])
    if not entry:
        return None
    return entry.entry_dn


def _get_group_dn(conn: Connection, cfg: ADConfig, group_name: str) -> Optional[str]:
    group_name = _safe_str(group_name)
    if not group_name:
        return None

    base_dn = _resolve_base_dn(conn, cfg)

    search_filter = (
        f"(&"
        f"(objectClass=group)"
        f"(|"
        f"(cn={_escape_ldap_filter_value(group_name)})"
        f"(name={_escape_ldap_filter_value(group_name)})"
        f"(sAMAccountName={_escape_ldap_filter_value(group_name)})"
        f")"
        f")"
    )
    entry = _search_one(conn, base_dn, search_filter, attributes=["distinguishedName"])
    if not entry:
        return None
    return entry.entry_dn


def _entry_exists_by_dn(conn: Connection, dn: str, object_class: Optional[str] = None) -> bool:
    dn = _safe_str(dn)
    if not dn:
        return False

    search_filter = "(objectClass=*)" if not object_class else f"(objectClass={object_class})"
    ok = conn.search(
        search_base=dn,
        search_filter=search_filter,
        search_scope=BASE,
        attributes=["distinguishedName"],
        size_limit=1,
    )
    return bool(ok and conn.entries)


# -----------------------------
# User reports / LDAP reads
# -----------------------------

def list_users(limit: int = 200) -> List[dict]:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        conn.search(
            base,
            "(&(objectClass=user)(!(objectClass=computer)))",
            search_scope=SUBTREE,
            attributes=["sAMAccountName", "displayName", "mail", "employeeID"],
            size_limit=limit,
        )
        return [
            {
                "username": _entry_str(e, "sAMAccountName"),
                "display_name": _entry_str(e, "displayName", "cn", "name"),
                "email": _entry_str(e, "mail"),
                "hr_id": _entry_str(e, "employeeID"),
            }
            for e in conn.entries
        ]
    finally:
        conn.unbind()

def _is_account_expired(account_expires_value) -> bool:
    raw = _filetime_to_int(account_expires_value)

    if raw in (0, 9223372036854775807, 9223372036854775808):
        return False

    expires_dt = _windows_filetime_to_datetime(raw)
    if not expires_dt:
        return False

    return expires_dt <= _now_utc()


def _ldap_dt_to_datetime(value):
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=timezone.utc)
        return value

    text = _safe_str(value)
    if not text:
        return None

    formats = [
        "%Y%m%d%H%M%S.0Z",
        "%Y%m%d%H%M%SZ",
        "%Y%m%d%H%M%S.%fZ",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except Exception:
            continue

    return None


def list_users_detailed(limit: int = 2000) -> list[dict]:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)

        conn.search(
            base,
            "(&(objectClass=user)(!(objectClass=computer)))",
            search_scope=SUBTREE,
            attributes=[
                "sAMAccountName",
                "displayName",
                "cn",
                "name",
                "mail",
                "employeeID",
                "department",
                "manager",
                "distinguishedName",
                "userAccountControl",
                "lockoutTime",
                "accountExpires",
                "lastLogonTimestamp",
                "lastLogon",
                "whenCreated",
                "pwdLastSet",
            ],
            size_limit=limit,
        )

        users = []
        for e in conn.entries:
            user_account_control = _filetime_to_int(_entry_value(e, "userAccountControl", default=0))
            lockout_time = _filetime_to_int(_entry_value(e, "lockoutTime", default=0))
            account_expires = _filetime_to_int(_entry_value(e, "accountExpires", default=0))

            last_logon_timestamp = _filetime_to_int(_entry_value(e, "lastLogonTimestamp", default=0))
            last_logon = _filetime_to_int(_entry_value(e, "lastLogon", default=0))

            effective_last_logon_raw = 0
            last_logon_source = ""

            if last_logon_timestamp > 0:
                effective_last_logon_raw = last_logon_timestamp
                last_logon_source = "lastLogonTimestamp"
            elif last_logon > 0:
                effective_last_logon_raw = last_logon
                last_logon_source = "lastLogon"

            effective_last_logon_dt = _windows_filetime_to_datetime(effective_last_logon_raw)
            account_expires_dt = _windows_filetime_to_datetime(account_expires)
            when_created_dt = _ldap_dt_to_datetime(_entry_value(e, "whenCreated"))

            manager_dn = _entry_str(e, "manager")
            dn = _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", ""))

            users.append(
                {
                    "username": _entry_str(e, "sAMAccountName"),
                    "display_name": _entry_str(e, "displayName", "cn", "name"),
                    "email": _entry_str(e, "mail"),
                    "hr_id": _entry_str(e, "employeeID"),
                    "department": _entry_str(e, "department"),
                    "manager": manager_dn,
                    "manager_name": _manager_display_from_dn(manager_dn),
                    "dn": dn,
                    "user_account_control": user_account_control,
                    "lockout_time": lockout_time,
                    "account_expires": account_expires,
                    "account_expires_display": _datetime_to_display(account_expires_dt),
                    "last_logon_timestamp": last_logon_timestamp,
                    "last_logon": last_logon,
                    "effective_last_logon": effective_last_logon_raw,
                    "effective_last_logon_display": _datetime_to_display(effective_last_logon_dt),
                    "last_logon_source": last_logon_source,
                    "when_created_dt": when_created_dt,
                    "when_created": _datetime_to_display(when_created_dt),
                    "pwd_last_set": _filetime_to_int(_entry_value(e, "pwdLastSet", default=0)),
                    "is_disabled": bool(user_account_control & 2),
                    "is_locked": lockout_time > 0,
                    "is_expired": _is_account_expired(account_expires),
                }
            )

        return users
    finally:
        conn.unbind()


def get_user_report_data(report_type: str = "all") -> list[dict]:
    users = list_users_detailed()
    report_type = _safe_lower(report_type) or "all"

    if report_type == "all":
        return users

    if report_type == "empty_attributes":
        return [
            u for u in users
            if not u.get("display_name") or not u.get("email") or not u.get("hr_id")
        ]

    if report_type == "without_managers":
        return [u for u in users if not u.get("manager")]

    if report_type == "duplicate_attributes":
        seen = {}
        duplicates = []

        for u in users:
            email = _safe_lower(u.get("email"))
            if not email:
                continue
            if email in seen:
                duplicates.append(u)
                if seen[email] not in duplicates:
                    duplicates.append(seen[email])
            else:
                seen[email] = u

        unique = []
        seen_usernames = set()
        for u in duplicates:
            username = _safe_str(u.get("username"))
            if username and username not in seen_usernames:
                seen_usernames.add(username)
                unique.append(u)
        return unique

    if report_type == "without_email":
        return [u for u in users if not _safe_str(u.get("email"))]

    if report_type == "without_hr_id":
        return [u for u in users if not _safe_str(u.get("hr_id"))]

    if report_type == "enabled":
        return [u for u in users if not u.get("is_disabled")]

    if report_type == "recently_created":
        threshold = _now_utc() - timedelta(days=30)
        return [
            u for u in users
            if u.get("when_created_dt") and u.get("when_created_dt") >= threshold
        ]

    if report_type == "disabled":
        return [u for u in users if u.get("is_disabled")]

    if report_type == "locked":
        return [u for u in users if u.get("is_locked")]

    if report_type == "expired":
        return [u for u in users if u.get("is_expired")]

    if report_type == "inactive":
        threshold = _now_utc() - timedelta(days=90)
        out = []
        for u in users:
            raw = _filetime_to_int(u.get("effective_last_logon"))
            dt = _windows_filetime_to_datetime(raw)
            if not dt or dt < threshold:
                out.append(u)
        return out

    if report_type == "real_last_logon":
        return [u for u in users if _filetime_to_int(u.get("effective_last_logon")) > 0]

    if report_type == "recently_logged_on":
        return sorted(
            [u for u in users if _filetime_to_int(u.get("effective_last_logon")) > 0],
            key=lambda x: _filetime_to_int(x.get("effective_last_logon")),
            reverse=True,
        )[:200]

    return users


def _group_type_flags(group_type_value) -> tuple[bool, str]:
    try:
        raw = int(group_type_value)
    except Exception:
        raw = 0

    unsigned_raw = raw & 0xFFFFFFFF
    is_security = bool(unsigned_raw & 0x80000000)

    if is_security:
        return True, "Security"
    return False, "Distribution"


def list_groups_detailed(limit: int = 5000) -> list[dict]:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        conn.search(
            base,
            "(objectClass=group)",
            search_scope=SUBTREE,
            attributes=[
                "cn",
                "name",
                "description",
                "distinguishedName",
                "sAMAccountName",
                "groupType",
                "member",
                "whenCreated",
                "whenChanged",
            ],
            size_limit=limit,
        )

        groups = []
        for e in conn.entries:
            member_dns = _entry_list(e, "member")
            member_count = len(member_dns)

            is_security, group_type_label = _group_type_flags(_entry_value(e, "groupType", default=0))

            when_created_dt = _ldap_dt_to_datetime(_entry_value(e, "whenCreated"))
            when_changed_dt = _ldap_dt_to_datetime(_entry_value(e, "whenChanged"))

            nested_count = sum(1 for dn in member_dns if dn.lower().startswith("cn="))
            member_preview = ", ".join(member_dns[:5])

            groups.append(
                {
                    "name": _entry_str(e, "cn", "name", "sAMAccountName"),
                    "description": _entry_str(e, "description"),
                    "dn": _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", "")),
                    "sam": _entry_str(e, "sAMAccountName"),
                    "group_type_raw": _entry_value(e, "groupType", default=0),
                    "group_type_label": group_type_label,
                    "is_security": is_security,
                    "member_count": member_count,
                    "member_dns": member_dns,
                    "member_preview": member_preview,
                    "nested_member_count": nested_count,
                    "when_created_dt": when_created_dt,
                    "when_changed_dt": when_changed_dt,
                    "when_created_display": _datetime_to_display(when_created_dt),
                    "when_changed_display": _datetime_to_display(when_changed_dt),
                }
            )

        return groups
    finally:
        conn.unbind()


def get_group_report_data(report_type: str = "all") -> list[dict]:
    groups = list_groups_detailed()
    report_type = _safe_lower(report_type) or "all"

    if report_type == "all":
        return groups

    if report_type == "with_members":
        return [g for g in groups if int(g.get("member_count") or 0) > 0]

    if report_type == "detailed_members":
        return [g for g in groups if int(g.get("member_count") or 0) > 0]

    if report_type == "without_members":
        return [g for g in groups if int(g.get("member_count") or 0) == 0]

    if report_type == "nested_groups":
        return [g for g in groups if int(g.get("nested_member_count") or 0) > 0]

    if report_type == "security":
        return [g for g in groups if g.get("is_security")]

    if report_type == "distribution":
        return [g for g in groups if not g.get("is_security")]

    if report_type == "without_description":
        return [g for g in groups if not _safe_str(g.get("description"))]

    if report_type == "large_groups":
        return [g for g in groups if int(g.get("member_count") or 0) >= 50]

    if report_type == "recently_created":
        threshold = _now_utc() - timedelta(days=30)
        return [
            g for g in groups
            if g.get("when_created_dt") and g.get("when_created_dt") >= threshold
        ]

    if report_type == "recently_modified":
        threshold = _now_utc() - timedelta(days=30)
        return [
            g for g in groups
            if g.get("when_changed_dt") and g.get("when_changed_dt") >= threshold
        ]

    return groups


# -----------------------------
# Computers / Groups / OUs LDAP reads
# -----------------------------

def list_computers(limit: int = 200) -> List[dict]:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        conn.search(
            base,
            "(objectClass=computer)",
            search_scope=SUBTREE,
            attributes=["sAMAccountName", "cn", "name", "distinguishedName", "description"],
            size_limit=limit,
        )
        return [
            {
                "name": _entry_str(e, "cn", "name", "sAMAccountName"),
                "sam": _entry_str(e, "sAMAccountName"),
                "dn": _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", "")),
                "description": _entry_str(e, "description"),
            }
            for e in conn.entries
        ]
    finally:
        conn.unbind()


def list_groups(limit: int = 5000) -> List[dict]:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        conn.search(
            base,
            "(objectClass=group)",
            search_scope=SUBTREE,
            attributes=["cn", "name", "description", "distinguishedName", "sAMAccountName"],
            size_limit=limit,
        )
        return [
            {
                "name": _entry_str(e, "cn", "name", "sAMAccountName"),
                "description": _entry_str(e, "description"),
                "dn": _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", "")),
                "sam": _entry_str(e, "sAMAccountName"),
            }
            for e in conn.entries
        ]
    finally:
        conn.unbind()


def list_ous(limit: int = 2000) -> List[dict]:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        conn.search(
            base,
            "(|(objectClass=organizationalUnit)(objectClass=container))",
            search_scope=SUBTREE,
            attributes=["ou", "cn", "name", "distinguishedName", "description"],
            size_limit=limit,
        )

        out = []
        for e in conn.entries:
            dn = _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", ""))
            name = _entry_str(e, "ou", "cn", "name") or dn
            out.append({
                "ou": name,
                "name": name,
                "dn": dn,
                "description": _entry_str(e, "description"),
            })

        out.sort(key=lambda x: _safe_lower(x.get("ou")))
        return out
    finally:
        conn.unbind()


def get_ad_counts() -> dict:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)

        def _count(filter_str: str) -> int:
            entries = conn.extend.standard.paged_search(
                search_base=base,
                search_filter=filter_str,
                search_scope=SUBTREE,
                attributes=["distinguishedName"],
                paged_size=1000,
                generator=False,
            )
            return len(entries)

        return {
            "users": _count("(&(objectClass=user)(!(objectClass=computer)))"),
            "computers": _count("(objectClass=computer)"),
            "groups": _count("(objectClass=group)"),
            "ous": _count("(objectClass=organizationalUnit)"),
        }
    finally:
        conn.unbind()

def _parent_dn_from_dn(dn: str) -> str:
    dn = _safe_str(dn)
    if not dn or "," not in dn:
        return ""
    return dn.split(",", 1)[1].strip()


def _ou_display_from_dn(dn: str) -> str:
    dn = _safe_str(dn)
    if not dn:
        return "-"
    first = dn.split(",", 1)[0]
    if "=" in first:
        return first.split("=", 1)[1]
    return dn


def _paged_entries(conn: Connection, base: str, search_filter: str, attributes: list[str]) -> list[dict]:
    results = conn.extend.standard.paged_search(
        search_base=base,
        search_filter=search_filter,
        search_scope=SUBTREE,
        attributes=attributes,
        paged_size=1000,
        generator=False,
    )
    return [item for item in (results or []) if item.get("type") == "searchResEntry"]


def _dash_attr(item: dict, attr_name: str, default=None):
    attrs = item.get("attributes") or {}
    value = attrs.get(attr_name, default)
    return value if value is not None else default


def _dash_str(item: dict, attr_name: str, default="") -> str:
    value = _dash_attr(item, attr_name, default)
    if isinstance(value, list):
        value = value[0] if value else default
    return _safe_str(value)


def _dash_list(item: dict, attr_name: str) -> list[str]:
    value = _dash_attr(item, attr_name, [])
    if value is None:
        return []
    if isinstance(value, list):
        return [_safe_str(v) for v in value if _safe_str(v)]
    text = _safe_str(value)
    return [text] if text else []


def _dash_dn(item: dict) -> str:
    return _safe_str(item.get("dn") or _dash_str(item, "distinguishedName"))


def _dash_when_created(item: dict):
    return _ldap_dt_to_datetime(_dash_attr(item, "whenCreated"))


def _dash_filetime_dt(item: dict, attr_name: str):
    return _windows_filetime_to_datetime(_filetime_to_int(_dash_attr(item, attr_name, 0)))


def get_dashboard_data(inactive_days: int = 90, top_limit: int = 5) -> dict:
    """Fast dashboard data using four paged LDAP searches."""
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        threshold = _now_utc() - timedelta(days=inactive_days)

        users = _paged_entries(conn, base, "(&(objectClass=user)(!(objectClass=computer)))", [
            "sAMAccountName", "displayName", "cn", "name", "distinguishedName",
            "userAccountControl", "lockoutTime", "lastLogonTimestamp", "whenCreated",
        ])
        computers = _paged_entries(conn, base, "(objectClass=computer)", [
            "sAMAccountName", "cn", "name", "distinguishedName", "userAccountControl",
            "lastLogonTimestamp", "operatingSystem", "whenCreated",
        ])
        groups = _paged_entries(conn, base, "(objectClass=group)", [
            "cn", "name", "sAMAccountName", "distinguishedName", "member", "whenCreated",
        ])
        ous = _paged_entries(conn, base, "(objectClass=organizationalUnit)", [
            "ou", "name", "distinguishedName", "whenCreated",
        ])

        disabled_users = locked_users = inactive_users = password_never_expires = 0
        recent_users = []
        for item in users:
            uac = _filetime_to_int(_dash_attr(item, "userAccountControl", 0))
            lockout_time = _filetime_to_int(_dash_attr(item, "lockoutTime", 0))
            last_logon_dt = _dash_filetime_dt(item, "lastLogonTimestamp")
            when_created_dt = _dash_when_created(item)
            dn = _dash_dn(item)

            disabled_users += 1 if uac & 2 else 0
            locked_users += 1 if lockout_time > 0 else 0
            password_never_expires += 1 if uac & 0x10000 else 0
            inactive_users += 1 if (not last_logon_dt or last_logon_dt < threshold) else 0

            recent_users.append({
                "name": _dash_str(item, "displayName") or _dash_str(item, "cn") or _dash_str(item, "name") or _dash_str(item, "sAMAccountName"),
                "sam": _dash_str(item, "sAMAccountName"),
                "ou": _ou_display_from_dn(_parent_dn_from_dn(dn)),
                "enabled": not bool(uac & 2),
                "when_created": _datetime_to_display(when_created_dt),
                "_created_dt": when_created_dt or datetime.min.replace(tzinfo=timezone.utc),
            })

        disabled_computers = inactive_computers = domain_controllers = 0
        os_counts = {}
        recent_computers = []
        for item in computers:
            uac = _filetime_to_int(_dash_attr(item, "userAccountControl", 0))
            last_logon_dt = _dash_filetime_dt(item, "lastLogonTimestamp")
            when_created_dt = _dash_when_created(item)
            dn = _dash_dn(item)
            os_name = _dash_str(item, "operatingSystem") or "Unknown"
            os_counts[os_name] = os_counts.get(os_name, 0) + 1

            disabled_computers += 1 if uac & 2 else 0
            inactive_computers += 1 if (not last_logon_dt or last_logon_dt < threshold) else 0
            domain_controllers += 1 if "ou=domain controllers" in _safe_lower(dn) else 0

            recent_computers.append({
                "name": _dash_str(item, "cn") or _dash_str(item, "name") or _dash_str(item, "sAMAccountName"),
                "sam": _dash_str(item, "sAMAccountName"),
                "ou": _ou_display_from_dn(_parent_dn_from_dn(dn)),
                "os": os_name,
                "enabled": not bool(uac & 2),
                "when_created": _datetime_to_display(when_created_dt),
                "_created_dt": when_created_dt or datetime.min.replace(tzinfo=timezone.utc),
            })

        top_groups = []
        empty_groups = 0
        for item in groups:
            count = len(_dash_list(item, "member"))
            empty_groups += 1 if count == 0 else 0
            top_groups.append({
                "name": _dash_str(item, "cn") or _dash_str(item, "name") or _dash_str(item, "sAMAccountName"),
                "count": count,
            })
        top_groups = sorted(top_groups, key=lambda x: x["count"], reverse=True)[:top_limit]

        ou_name_by_dn = {}
        child_counts_by_parent = {}
        for item in ous:
            dn = _dash_dn(item)
            ou_name_by_dn[_safe_lower(dn)] = _dash_str(item, "ou") or _dash_str(item, "name") or _ou_display_from_dn(dn)
        for collection in (users, computers, groups, ous):
            for item in collection:
                parent_dn = _parent_dn_from_dn(_dash_dn(item))
                if parent_dn:
                    key = _safe_lower(parent_dn)
                    child_counts_by_parent[key] = child_counts_by_parent.get(key, 0) + 1
        top_ous = sorted(
            [{"name": name, "count": child_counts_by_parent.get(dn_key, 0)} for dn_key, name in ou_name_by_dn.items()],
            key=lambda x: x["count"],
            reverse=True,
        )[:top_limit]

        recent_users = sorted(recent_users, key=lambda x: x["_created_dt"], reverse=True)[:top_limit]
        recent_computers = sorted(recent_computers, key=lambda x: x["_created_dt"], reverse=True)[:top_limit]
        for row in recent_users + recent_computers:
            row.pop("_created_dt", None)

        top_os = sorted(os_counts.items(), key=lambda x: x[1], reverse=True)[:top_limit]

        return {
            "ad_counts": {"users": len(users), "computers": len(computers), "groups": len(groups), "ous": len(ous)},
            "user_status": {
                "enabled": max(len(users) - disabled_users, 0),
                "disabled": disabled_users,
                "locked": locked_users,
                "inactive": inactive_users,
                "password_never_expires": password_never_expires,
            },
            "computer_status": {
                "enabled": max(len(computers) - disabled_computers, 0),
                "disabled": disabled_computers,
                "inactive": inactive_computers,
                "domain_controllers": domain_controllers,
            },
            "top_ous": top_ous,
            "top_groups": top_groups,
            "computer_os": {"labels": [x[0] for x in top_os], "data": [x[1] for x in top_os]},
            "security_alerts": [
                {"name": "Disabled users", "count": disabled_users, "level": "danger" if disabled_users else "success", "status": "Review" if disabled_users else "OK"},
                {"name": "Locked users", "count": locked_users, "level": "warning" if locked_users else "success", "status": "Unlock" if locked_users else "OK"},
                {"name": f"Inactive users ({inactive_days}+ days)", "count": inactive_users, "level": "warning" if inactive_users else "success", "status": "Review" if inactive_users else "OK"},
                {"name": "Password never expires", "count": password_never_expires, "level": "warning" if password_never_expires else "success", "status": "Risk" if password_never_expires else "OK"},
                {"name": "Disabled computers", "count": disabled_computers, "level": "danger" if disabled_computers else "success", "status": "Review" if disabled_computers else "OK"},
                {"name": "Empty groups", "count": empty_groups, "level": "warning" if empty_groups else "success", "status": "Cleanup" if empty_groups else "OK"},
            ],
            "recent_users": recent_users,
            "recent_computers": recent_computers,
        }
    finally:
        conn.unbind()




def _computer_to_report_dict(conn: Connection, entry) -> dict:
    """Normalize a computer LDAP entry into a template-safe report dictionary."""
    user_account_control = _filetime_to_int(_entry_value(entry, "userAccountControl", default=0))
    last_logon_timestamp = _filetime_to_int(_entry_value(entry, "lastLogonTimestamp", default=0))
    last_logon = _filetime_to_int(_entry_value(entry, "lastLogon", default=0))

    effective_last_logon_raw = 0
    last_logon_source = ""

    if last_logon_timestamp > 0:
        effective_last_logon_raw = last_logon_timestamp
        last_logon_source = "lastLogonTimestamp"
    elif last_logon > 0:
        effective_last_logon_raw = last_logon
        last_logon_source = "lastLogon"

    effective_last_logon_dt = _windows_filetime_to_datetime(effective_last_logon_raw)
    when_created_dt = _ldap_dt_to_datetime(_entry_value(entry, "whenCreated"))
    when_changed_dt = _ldap_dt_to_datetime(_entry_value(entry, "whenChanged"))

    dn = _entry_str(entry, "distinguishedName") or _safe_str(getattr(entry, "entry_dn", ""))
    os_name = _entry_str(entry, "operatingSystem")
    recovery_key = ""
    recovery_key_created = ""

    # BitLocker recovery passwords are usually stored as child msFVE-RecoveryInformation objects.
    # This query is best-effort and safely returns empty values if the account has no permission
    # or if no BitLocker recovery objects exist below the computer object.
    if dn:
        try:
            conn.search(
                search_base=dn,
                search_filter="(objectClass=msFVE-RecoveryInformation)",
                search_scope=SUBTREE,
                attributes=["msFVE-RecoveryPassword", "whenCreated", "distinguishedName"],
                size_limit=5,
            )
            recovery_entries = list(conn.entries or [])
            recovery_entries.sort(
                key=lambda x: _safe_str(_entry_value(x, "whenCreated")),
                reverse=True,
            )
            if recovery_entries:
                latest = recovery_entries[0]
                recovery_key = _entry_str(latest, "msFVE-RecoveryPassword")
                recovery_key_created_dt = _ldap_dt_to_datetime(_entry_value(latest, "whenCreated"))
                recovery_key_created = _datetime_to_display(recovery_key_created_dt)
        except Exception:
            recovery_key = ""
            recovery_key_created = ""

    return {
        "name": _entry_str(entry, "cn", "name", "sAMAccountName"),
        "sam": _entry_str(entry, "sAMAccountName"),
        "dn": dn,
        "description": _entry_str(entry, "description"),
        "dns_host_name": _entry_str(entry, "dNSHostName"),
        "os": os_name,
        "operating_system": os_name,
        "os_version": _entry_str(entry, "operatingSystemVersion"),
        "user_account_control": user_account_control,
        "is_disabled": bool(user_account_control & 2),
        "is_active": not bool(user_account_control & 2),
        "last_logon": effective_last_logon_raw,
        "last_logon_display": _datetime_to_display(effective_last_logon_dt),
        "last_logon_source": last_logon_source,
        "when_created_dt": when_created_dt,
        "when_changed_dt": when_changed_dt,
        "when_created": _datetime_to_display(when_created_dt),
        "when_changed": _datetime_to_display(when_changed_dt),
        "recovery_key": recovery_key,
        "bitlocker_recovery_key": recovery_key,
        "recovery_key_created": recovery_key_created,
        "has_bitlocker_key": bool(recovery_key),
    }


def list_computers_detailed(limit: int = 5000) -> list[dict]:
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        conn.search(
            base,
            "(objectClass=computer)",
            search_scope=SUBTREE,
            attributes=[
                "sAMAccountName",
                "cn",
                "name",
                "distinguishedName",
                "description",
                "dNSHostName",
                "operatingSystem",
                "operatingSystemVersion",
                "userAccountControl",
                "lastLogonTimestamp",
                "lastLogon",
                "whenCreated",
                "whenChanged",
            ],
            size_limit=limit,
        )

        computers = [_computer_to_report_dict(conn, e) for e in conn.entries]
        computers.sort(key=lambda x: _safe_lower(x.get("name")))
        return computers
    finally:
        conn.unbind()


def get_computer_report_data(report_type: str = "all") -> list[dict]:
    computers = list_computers_detailed()
    report_type = _safe_lower(report_type) or "all"

    if report_type == "all":
        return computers

    if report_type == "os_based":
        return [c for c in computers if _safe_str(c.get("os"))]

    if report_type == "workstations":
        out = []
        for c in computers:
            os_name = _safe_lower(c.get("os"))
            dn = _safe_lower(c.get("dn"))
            is_server = "server" in os_name or "domain controllers" in dn
            if os_name and not is_server:
                out.append(c)
        return out

    if report_type == "inactive":
        threshold = _now_utc() - timedelta(days=90)
        out = []
        for c in computers:
            raw = _filetime_to_int(c.get("last_logon"))
            dt = _windows_filetime_to_datetime(raw)
            if not dt or dt < threshold:
                out.append(c)
        return out

    if report_type == "active":
        return [c for c in computers if not c.get("is_disabled")]

    if report_type == "disabled":
        return [c for c in computers if c.get("is_disabled")]

    if report_type == "bitlocker_keys":
        return [c for c in computers if c.get("has_bitlocker_key")]

    if report_type == "bitlocker_enabled":
        return [c for c in computers if c.get("has_bitlocker_key")]

    return computers


def _get_ou_protection_map_winrm(cfg: ADConfig) -> dict[str, bool]:
    """Best-effort OU protection lookup using AD PowerShell when WinRM is configured."""
    if not cfg.server_name or session is None:
        return {}

    server_esc = _ps_escape_single_quotes(cfg.server_name)
    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory
Get-ADOrganizationalUnit -Filter * -Server '{server_esc}' -Properties ProtectedFromAccidentalDeletion |
    Select-Object DistinguishedName,ProtectedFromAccidentalDeletion |
    ConvertTo-Json -Depth 3
"""
    try:
        data = _parse_json_output(_run_ps(ps))
        return {
            _safe_lower(item.get("DistinguishedName")): bool(item.get("ProtectedFromAccidentalDeletion"))
            for item in data
            if _safe_str(item.get("DistinguishedName"))
        }
    except Exception:
        return {}


def _count_immediate_children(conn: Connection, ou_dn: str, search_filter: str) -> int:
    try:
        conn.search(
            search_base=ou_dn,
            search_filter=search_filter,
            search_scope=LEVEL,
            attributes=["distinguishedName"],
            size_limit=0,
        )
        return len(conn.entries or [])
    except Exception:
        return 0


def _ou_entry_to_report_dict(entry, include_dates: bool = True) -> dict:
    dn = _entry_str(entry, "distinguishedName") or _safe_str(getattr(entry, "entry_dn", ""))
    when_created_dt = _ldap_dt_to_datetime(_entry_value(entry, "whenCreated")) if include_dates else None
    when_changed_dt = _ldap_dt_to_datetime(_entry_value(entry, "whenChanged")) if include_dates else None
    gp_link = _entry_str(entry, "gPLink")

    return {
        "name": _entry_str(entry, "ou", "cn", "name") or dn,
        "ou": _entry_str(entry, "ou", "cn", "name") or dn,
        "description": _entry_str(entry, "description"),
        "dn": dn,
        "when_created_dt": when_created_dt,
        "when_changed_dt": when_changed_dt,
        "when_created": _datetime_to_display(when_created_dt),
        "when_changed": _datetime_to_display(when_changed_dt),
        "gp_link": gp_link,
        "users_count": 0,
        "computers_count": 0,
        "groups_count": 0,
        "child_ous_count": 0,
        "total_count": 0,
        "is_empty": False,
        "is_protected": False,
        "is_unprotected": False,
        "has_gpo_link": bool(gp_link),
    }


def _search_ou_entries(conn: Connection, base: str, search_filter: str = "(objectClass=organizationalUnit)", limit: int = 5000):
    conn.search(
        base,
        search_filter,
        search_scope=SUBTREE,
        attributes=[
            "ou",
            "cn",
            "name",
            "description",
            "distinguishedName",
            "whenCreated",
            "whenChanged",
            "gPLink",
        ],
        size_limit=limit,
    )
    return list(conn.entries or [])


def _add_ou_counts(conn: Connection, ou: dict) -> dict:
    dn = _safe_str(ou.get("dn"))
    if not dn:
        return ou

    users_count = _count_immediate_children(
        conn,
        dn,
        "(&(objectClass=user)(!(objectClass=computer)))",
    )
    computers_count = _count_immediate_children(conn, dn, "(objectClass=computer)")
    groups_count = _count_immediate_children(conn, dn, "(objectClass=group)")
    child_ous_count = _count_immediate_children(conn, dn, "(objectClass=organizationalUnit)")
    total_count = users_count + computers_count + groups_count + child_ous_count

    ou.update(
        {
            "users_count": users_count,
            "computers_count": computers_count,
            "groups_count": groups_count,
            "child_ous_count": child_ous_count,
            "total_count": total_count,
            "is_empty": total_count == 0,
        }
    )
    return ou


def _get_ous_with_counts(conn: Connection, base: str, limit: int = 5000) -> list[dict]:
    entries = _search_ou_entries(conn, base, limit=limit)
    ous = [_ou_entry_to_report_dict(e) for e in entries]

    for ou in ous:
        _add_ou_counts(conn, ou)

    ous.sort(key=lambda x: _safe_lower(x.get("name")))
    return ous


def list_ous_detailed(limit: int = 5000, include_counts: bool = False, include_protection: bool = False) -> list[dict]:
    """
    Fast OU list for reports.

    Important:
    - Counts are expensive because they require child searches per OU.
    - Protection requires WinRM/AD PowerShell.
    - Both are optional so normal reports load quickly.
    """
    cfg = _load_config()
    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        entries = _search_ou_entries(conn, base, limit=limit)
        ous = [_ou_entry_to_report_dict(e) for e in entries]

        if include_counts:
            for ou in ous:
                _add_ou_counts(conn, ou)

        if include_protection:
            protection_map = _get_ou_protection_map_winrm(cfg)
            for ou in ous:
                protected = protection_map.get(_safe_lower(ou.get("dn")), False)
                ou["is_protected"] = protected
                ou["is_unprotected"] = not protected

        ous.sort(key=lambda x: _safe_lower(x.get("name")))
        return ous
    finally:
        conn.unbind()


def get_ou_report_data(report_type: str = "all") -> list[dict]:
    """
    Return OU report data using the fastest possible path for each report.

    Fast reports do one LDAP search:
    - all
    - recently_created
    - recently_modified
    - gpo_linked

    Expensive reports only do expensive work when selected:
    - counts / empty: per-OU immediate child counts
    - protected / unprotected: one WinRM PowerShell protection lookup
    """
    report_type = _safe_lower(report_type) or "all"

    cfg = _load_config()
    conn = _connect(cfg)

    try:
        base = _resolve_base_dn(conn, cfg)

        if report_type == "gpo_linked":
            entries = _search_ou_entries(
                conn,
                base,
                search_filter="(&(objectClass=organizationalUnit)(gPLink=*))",
            )
            ous = [_ou_entry_to_report_dict(e) for e in entries]
            ous.sort(key=lambda x: _safe_lower(x.get("name")))
            return ous

        entries = _search_ou_entries(conn, base)
        ous = [_ou_entry_to_report_dict(e) for e in entries]

        if report_type == "all":
            ous.sort(key=lambda x: _safe_lower(x.get("name")))
            return ous

        if report_type == "recently_created":
            threshold = _now_utc() - timedelta(days=30)
            out = [
                ou for ou in ous
                if ou.get("when_created_dt") and ou.get("when_created_dt") >= threshold
            ]
            out.sort(key=lambda x: _safe_lower(x.get("name")))
            return out

        if report_type == "recently_modified":
            threshold = _now_utc() - timedelta(days=30)
            out = [
                ou for ou in ous
                if ou.get("when_changed_dt") and ou.get("when_changed_dt") >= threshold
            ]
            out.sort(key=lambda x: _safe_lower(x.get("name")))
            return out

        if report_type in {"counts", "empty"}:
            for ou in ous:
                _add_ou_counts(conn, ou)

            if report_type == "empty":
                ous = [ou for ou in ous if ou.get("is_empty")]

            ous.sort(key=lambda x: _safe_lower(x.get("name")))
            return ous

        if report_type in {"protected", "unprotected"}:
            protection_map = _get_ou_protection_map_winrm(cfg)
            for ou in ous:
                protected = protection_map.get(_safe_lower(ou.get("dn")), False)
                ou["is_protected"] = protected
                ou["is_unprotected"] = not protected

            if report_type == "protected":
                ous = [ou for ou in ous if ou.get("is_protected")]
            else:
                ous = [ou for ou in ous if not ou.get("is_protected")]

            ous.sort(key=lambda x: _safe_lower(x.get("name")))
            return ous

        ous.sort(key=lambda x: _safe_lower(x.get("name")))
        return ous
    finally:
        conn.unbind()


def get_ou_details(ou_dn: str) -> dict:
    ou_dn = _safe_str(ou_dn)
    if not ou_dn:
        raise ADServiceError("OU DN is required.")

    cfg = _load_config()
    conn = _connect(cfg)
    try:
        conn.search(
            search_base=ou_dn,
            search_filter="(objectClass=organizationalUnit)",
            search_scope=BASE,
            attributes=[
                "ou",
                "cn",
                "name",
                "description",
                "distinguishedName",
                "whenCreated",
                "whenChanged",
                "gPLink",
            ],
            size_limit=1,
        )

        if not conn.entries:
            raise ADServiceError(f"OU not found: {ou_dn}")

        e = conn.entries[0]
        when_created_dt = _ldap_dt_to_datetime(_entry_value(e, "whenCreated"))
        when_changed_dt = _ldap_dt_to_datetime(_entry_value(e, "whenChanged"))

        return {
            "name": _entry_str(e, "ou", "cn", "name") or ou_dn,
            "ou": _entry_str(e, "ou", "cn", "name") or ou_dn,
            "description": _entry_str(e, "description"),
            "dn": _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", "")),
            "when_created": _datetime_to_display(when_created_dt),
            "when_changed": _datetime_to_display(when_changed_dt),
            "gp_link": _entry_str(e, "gPLink"),
        }
    finally:
        conn.unbind()



# -----------------------------
# Friendly-name resolvers
# -----------------------------

def resolve_ou_name_to_dn(ou_name: str) -> str:
    target = _safe_lower(ou_name)
    if not target:
        raise ADServiceError("OU name is required.")

    ous = list_ous()
    matches = [ou for ou in ous if _safe_lower(ou.get("ou")) == target]

    if not matches:
        raise ADServiceError(f"OU '{ou_name}' not found.")

    if len(matches) > 1:
        matched_dns = ", ".join(m["dn"] for m in matches if _safe_str(m.get("dn")))
        raise ADServiceError(f"OU '{ou_name}' is ambiguous. Matches: {matched_dns}")

    return matches[0]["dn"]


def resolve_user_username_to_dn(conn: Connection, cfg: ADConfig, username: str) -> str:
    username = _safe_str(username)
    if not username:
        raise ADServiceError("Username is required.")

    dn = _get_user_dn(conn, cfg, username)
    if not dn:
        raise ADServiceError(f"User '{username}' not found.")
    return dn


def resolve_computer_name_to_dn(conn: Connection, cfg: ADConfig, computer_name: str) -> str:
    computer_name = _safe_str(computer_name)
    if not computer_name:
        raise ADServiceError("Computer name is required.")

    dn = _get_computer_dn(conn, cfg, computer_name)
    if not dn:
        raise ADServiceError(f"Computer '{computer_name}' not found.")
    return dn


def resolve_group_name_to_dn(conn: Connection, cfg: ADConfig, group_name: str) -> str:
    group_name = _safe_str(group_name)
    if not group_name:
        raise ADServiceError("Group name is required.")

    base = _resolve_base_dn(conn, cfg)
    conn.search(
        base,
        (
            f"(&"
            f"(objectClass=group)"
            f"(|"
            f"(cn={_escape_ldap_filter_value(group_name)})"
            f"(name={_escape_ldap_filter_value(group_name)})"
            f"(sAMAccountName={_escape_ldap_filter_value(group_name)})"
            f")"
            f")"
        ),
        search_scope=SUBTREE,
        attributes=["distinguishedName"],
        size_limit=5,
    )

    matches = conn.entries or []
    if not matches:
        raise ADServiceError(f"Group '{group_name}' not found.")

    if len(matches) > 1:
        dns = ", ".join(e.entry_dn for e in matches)
        raise ADServiceError(f"Group '{group_name}' is ambiguous. Matches: {dns}")

    return matches[0].entry_dn


# -----------------------------
# Search endpoints for AJAX pickers
# -----------------------------

def search_users(query: str, limit: int = 15) -> list[dict]:
    cfg = _load_config()
    if cfg.safe_mode:
        return []

    query = _safe_str(query)
    if len(query) < 2:
        return []

    cache_key = ("users", query.lower(), limit)
    cached = _cache_get(cache_key)
    if cached is not None:
        return cached

    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        q = _escape_ldap_filter_value(query)

        search_filter = (
            f"(&"
            f"(objectClass=user)"
            f"(!(objectClass=computer))"
            f"(|"
            f"(sAMAccountName={q}*)"
            f"(displayName={q}*)"
            f"(name={q}*)"
            f"(cn={q}*)"
            f"(mail={q}*)"
            f")"
            f")"
        )

        conn.search(
            search_base=base,
            search_filter=search_filter,
            search_scope=SUBTREE,
            attributes=["displayName", "cn", "name", "sAMAccountName", "distinguishedName"],
            size_limit=limit,
        )

        results = []
        for e in conn.entries:
            results.append(
                {
                    "DisplayName": _entry_str(e, "displayName", "cn", "name"),
                    "SamAccountName": _entry_str(e, "sAMAccountName"),
                    "DistinguishedName": _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", "")),
                }
            )

        results.sort(key=lambda x: (
            0 if _safe_lower(x["SamAccountName"]).startswith(query.lower()) else 1,
            0 if _safe_lower(x["DisplayName"]).startswith(query.lower()) else 1,
            _safe_lower(x["SamAccountName"]),
        ))

        _cache_set(cache_key, results)
        return results
    finally:
        conn.unbind()


def search_computers(query: str, limit: int = 15) -> list[dict]:
    cfg = _load_config()
    if cfg.safe_mode:
        return []

    query = _safe_str(query)
    if len(query) < 2:
        return []

    cache_key = ("computers", query.lower(), limit)
    cached = _cache_get(cache_key)
    if cached is not None:
        return cached

    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        q = _escape_ldap_filter_value(query)
        q_sam = _escape_ldap_filter_value(_normalize_computer_sam(query))

        search_filter = (
            f"(&"
            f"(objectClass=computer)"
            f"(|"
            f"(cn={q}*)"
            f"(name={q}*)"
            f"(sAMAccountName={q_sam}*)"
            f")"
            f")"
        )

        conn.search(
            search_base=base,
            search_filter=search_filter,
            search_scope=SUBTREE,
            attributes=["name", "cn", "distinguishedName", "sAMAccountName"],
            size_limit=limit,
        )

        results = []
        for e in conn.entries:
            results.append(
                {
                    "Name": _entry_str(e, "name", "cn", "sAMAccountName"),
                    "SamAccountName": _entry_str(e, "sAMAccountName"),
                    "DistinguishedName": _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", "")),
                }
            )

        results.sort(key=lambda x: (
            0 if _safe_lower(x["Name"]).startswith(query.lower()) else 1,
            _safe_lower(x["Name"]),
        ))

        _cache_set(cache_key, results)
        return results
    finally:
        conn.unbind()


def search_groups(query: str, limit: int = 15) -> list[dict]:
    cfg = _load_config()
    if cfg.safe_mode:
        return []

    query = _safe_str(query)
    if len(query) < 2:
        return []

    cache_key = ("groups", query.lower(), limit)
    cached = _cache_get(cache_key)
    if cached is not None:
        return cached

    conn = _connect(cfg)
    try:
        base = _resolve_base_dn(conn, cfg)
        q = _escape_ldap_filter_value(query)

        search_filter = (
            f"(&"
            f"(objectClass=group)"
            f"(|"
            f"(cn={q}*)"
            f"(name={q}*)"
            f"(sAMAccountName={q}*)"
            f")"
            f")"
        )

        conn.search(
            search_base=base,
            search_filter=search_filter,
            search_scope=SUBTREE,
            attributes=["cn", "name", "sAMAccountName", "distinguishedName"],
            size_limit=limit,
        )

        results = []
        for e in conn.entries:
            results.append(
                {
                    "Name": _entry_str(e, "cn", "name", "sAMAccountName"),
                    "SamAccountName": _entry_str(e, "sAMAccountName"),
                    "DistinguishedName": _entry_str(e, "distinguishedName") or _safe_str(getattr(e, "entry_dn", "")),
                }
            )

        results.sort(key=lambda x: (
            0 if _safe_lower(x["Name"]).startswith(query.lower()) else 1,
            _safe_lower(x["Name"]),
        ))

        _cache_set(cache_key, results)
        return results
    finally:
        conn.unbind()


def get_group_members_for_ui(group_name: str) -> list[dict]:
    return get_group_members_for_update(group_name)


def search_directory_objects(query: str, limit: int = 15) -> list[dict]:
    query = _safe_str(query)
    if not query:
        return []

    results = []

    for user in search_users(query, limit=limit):
        label = user.get("DisplayName") or user.get("SamAccountName") or "Unknown User"
        sam = user.get("SamAccountName") or ""
        if sam and sam not in label:
            label = f"{label} ({sam})"

        results.append(
            {
                "label": label,
                "value": user.get("DistinguishedName", ""),
                "type": "user",
            }
        )

    for computer in search_computers(query, limit=limit):
        label = computer.get("Name") or computer.get("SamAccountName") or "Unknown Computer"
        results.append(
            {
                "label": label,
                "value": computer.get("DistinguishedName", ""),
                "type": "computer",
            }
        )

    for group in search_groups(query, limit=limit):
        label = group.get("Name") or group.get("SamAccountName") or "Unknown Group"
        results.append(
            {
                "label": label,
                "value": group.get("DistinguishedName", ""),
                "type": "group",
            }
        )

    seen = set()
    unique_results = []
    for item in results:
        key = (item["type"], item["value"])
        if key in seen or not item["value"]:
            continue
        seen.add(key)
        unique_results.append(item)

    unique_results.sort(key=lambda x: (x["type"], _safe_lower(x["label"])))
    return unique_results[:limit]


def list_group_ous(limit: int = 100) -> list[dict]:
    cfg = _load_config()
    if cfg.safe_mode:
        return []

    # Prefer LDAP for compatibility; fall back to scoped filtering if configured.
    ous = list_ous(limit=5000)
    groups_base = _safe_lower(cfg.groups_ou_dn)

    if not groups_base:
        return ous[:limit]

    filtered = [ou for ou in ous if _safe_lower(ou.get("dn")).endswith(groups_base)]
    return filtered[:limit] if filtered else ous[:limit]


# -----------------------------
# LDAP group membership helpers
# -----------------------------

def _get_group_members_dns(conn: Connection, group_dn: str) -> List[str]:
    ok = conn.search(
        search_base=group_dn,
        search_filter="(objectClass=group)",
        search_scope=BASE,
        attributes=["member"],
        size_limit=1,
    )
    if not ok or not conn.entries:
        return []

    entry = conn.entries[0]
    try:
        if "member" not in entry:
            return []
        return [str(v).strip() for v in entry.member.values if str(v).strip()]
    except Exception:
        return []


def get_group_members_for_update(group_name: str) -> list[dict]:
    cfg = _load_config()
    if cfg.safe_mode:
        return []

    conn = _connect(cfg)
    try:
        group_dn = _get_group_dn(conn, cfg, group_name)
        if not group_dn:
            raise ADServiceError(f"Group not found: {group_name}")

        member_dns = _get_group_members_dns(conn, group_dn)
        results: list[dict] = []

        for member_dn in member_dns:
            ok = conn.search(
                search_base=member_dn,
                search_filter="(objectClass=*)",
                search_scope=BASE,
                attributes=["objectClass", "displayName", "sAMAccountName", "cn", "name", "distinguishedName"],
                size_limit=1,
            )
            if not ok or not conn.entries:
                continue

            entry = conn.entries[0]
            object_classes = [str(v).lower() for v in _entry_list(entry, "objectClass")]

            if "computer" in object_classes:
                member_type = "computer"
                label = _entry_str(entry, "cn", "name", "sAMAccountName") or member_dn
            elif "group" in object_classes:
                member_type = "group"
                label = _entry_str(entry, "cn", "name", "sAMAccountName") or member_dn
            else:
                member_type = "user"
                display_name = _entry_str(entry, "displayName", "cn", "name")
                sam = _entry_str(entry, "sAMAccountName")
                label = f"{display_name} ({sam})" if display_name and sam else (display_name or sam or member_dn)

            results.append(
                {
                    "label": label,
                    "value": member_dn,
                    "type": member_type,
                }
            )

        results.sort(key=lambda x: (x["type"], _safe_lower(x["label"])))
        return results
    finally:
        conn.unbind()


def _set_group_owner_ldap(conn: Connection, group_dn: str, owner_dn: str) -> None:
    ok = conn.modify(group_dn, {"managedBy": [(MODIFY_REPLACE, [owner_dn])]})
    if not ok:
        raise ADServiceError(f"Failed to set group owner: {conn.result}")


def _add_members_to_group_ldap(conn: Connection, group_dn: str, member_dns: Iterable[str]) -> int:
    members = [str(m).strip() for m in (member_dns or []) if str(m).strip()]
    if not members:
        return 0

    existing = set(_get_group_members_dns(conn, group_dn))
    to_add = [m for m in members if m not in existing]

    added = 0
    for batch in _chunked(to_add, size=100):
        ok = conn.modify(group_dn, {"member": [(MODIFY_ADD, batch)]})
        if not ok:
            result = conn.result or {}
            desc = str(result.get("description", "")).lower()
            msg = str(result.get("message", ""))
            if "attributeorvalueexists" in desc:
                continue
            raise ADServiceError(f"Failed to add group members: {msg or result}")
        added += len(batch)

    return added


def _remove_members_from_group_ldap(conn: Connection, group_dn: str, member_dns: Iterable[str]) -> int:
    members = [str(m).strip() for m in (member_dns or []) if str(m).strip()]
    if not members:
        return 0

    existing = set(_get_group_members_dns(conn, group_dn))
    to_remove = [m for m in members if m in existing]

    removed = 0
    for batch in _chunked(to_remove, size=100):
        ok = conn.modify(group_dn, {"member": [(MODIFY_DELETE, batch)]})
        if not ok:
            raise ADServiceError(f"Failed to remove group members: {conn.result}")
        removed += len(batch)

    return removed


# -----------------------------
# WinRM helpers for group write operations
# -----------------------------

def _update_group_scope_category_winrm(
    cfg: ADConfig,
    group_dn: str,
    group_scope: str = "",
    group_category: str = "",
) -> None:
    normalized_scope = _normalize_group_scope(group_scope) if group_scope else ""
    normalized_category = _normalize_group_category(group_category) if group_category else ""

    if not normalized_scope and not normalized_category:
        return

    group_dn_esc = _ps_escape_single_quotes(group_dn)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    scope_line = ""
    category_line = ""

    if normalized_scope:
        scope_esc = _ps_escape_single_quotes(normalized_scope)
        scope_line = f"$params['GroupScope'] = '{scope_esc}'"

    if normalized_category:
        category_esc = _ps_escape_single_quotes(normalized_category)
        category_line = f"$params['GroupCategory'] = '{category_esc}'"

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$group = Get-ADGroup -Identity '{group_dn_esc}' -Server '{server_esc}' -ErrorAction Stop
$params = @{{
    Identity = $group.DistinguishedName
    Server   = '{server_esc}'
}}

{scope_line}
{category_line}

Set-ADGroup @params
Write-Output "UPDATED_SCOPE_CATEGORY=1"
"""
    out = _run_ps(ps)
    if "UPDATED_SCOPE_CATEGORY=1" not in out:
        raise ADServiceError("Failed to update group scope/category.")


def _protect_group_from_deletion_winrm(cfg: ADConfig, group_dn: str) -> None:
    group_dn_esc = _ps_escape_single_quotes(group_dn)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$group = Get-ADGroup -Identity '{group_dn_esc}' -Server '{server_esc}' -ErrorAction Stop
$de = [ADSI]("LDAP://" + $group.DistinguishedName)
$sd = $de.ObjectSecurity
$everyone = New-Object System.Security.Principal.NTAccount("Everyone")
$guid = [Guid]::Empty

$rule1 = New-Object System.DirectoryServices.ActiveDirectoryAccessRule(
    $everyone,
    [System.DirectoryServices.ActiveDirectoryRights]::Delete,
    [System.Security.AccessControl.AccessControlType]::Deny,
    $guid
)
$rule2 = New-Object System.DirectoryServices.ActiveDirectoryAccessRule(
    $everyone,
    [System.DirectoryServices.ActiveDirectoryRights]::DeleteTree,
    [System.Security.AccessControl.AccessControlType]::Deny,
    $guid
)

$sd.AddAccessRule($rule1)
$sd.AddAccessRule($rule2)
$de.ObjectSecurity = $sd
$de.CommitChanges()
Write-Output "PROTECT_DELETE=1"
"""
    out = _run_ps(ps)
    if "PROTECT_DELETE=1" not in out:
        raise ADServiceError("Group created, but failed to enable accidental deletion protection.")


# -----------------------------
# Group operations
# -----------------------------

def create_group(
    group_name: str,
    description: str = "",
    user_dns: Optional[Sequence[str]] = None,
    computer_dns: Optional[Sequence[str]] = None,
    nested_group_dns: Optional[Sequence[str]] = None,
    group_scope: str = "Global",
    group_category: str = "Security",
    owner_dn: str = "",
    ou_dn: str = "",
    protect_from_deletion: bool = False,
    copy_from_group_dn: str = "",
    skip_invalid_members: bool = True,
) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return

    allowed_scopes = {"Global", "DomainLocal", "Universal"}
    allowed_categories = {"Security", "Distribution"}

    if group_scope not in allowed_scopes:
        raise ADServiceError(f"Invalid group scope: {group_scope}")
    if group_category not in allowed_categories:
        raise ADServiceError(f"Invalid group category: {group_category}")

    group_name = _safe_str(group_name)
    description = _safe_str(description)
    owner_dn = _safe_str(owner_dn)
    target_ou = _safe_str(ou_dn) or _safe_str(cfg.groups_ou_dn)
    copy_from_group_dn = _safe_str(copy_from_group_dn)

    if not group_name:
        raise ADServiceError("Group name is required.")
    if not target_ou:
        raise ADServiceError("Target OU is required.")
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for group creation.")

    group_name_esc = _ps_escape_single_quotes(group_name)
    description_esc = _ps_escape_single_quotes(description)
    path_esc = _ps_escape_single_quotes(target_ou)
    server_esc = _ps_escape_single_quotes(cfg.server_name)
    scope_esc = _ps_escape_single_quotes(group_scope)
    category_esc = _ps_escape_single_quotes(group_category)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$existing = Get-ADGroup -LDAPFilter "(cn={group_name_esc})" -Server '{server_esc}' -ErrorAction SilentlyContinue
if ($existing) {{
    throw "Group already exists: $($existing.DistinguishedName)"
}}

$params = @{{
    Name           = '{group_name_esc}'
    SamAccountName = '{group_name_esc}'
    GroupScope     = '{scope_esc}'
    GroupCategory  = '{category_esc}'
    Path           = '{path_esc}'
    Server         = '{server_esc}'
}}

if ('{description_esc}') {{
    $params['Description'] = '{description_esc}'
}}

New-ADGroup @params

$group = Get-ADGroup -Identity '{group_name_esc}' -Server '{server_esc}' -Properties DistinguishedName
if (-not $group) {{
    throw "Group creation succeeded but could not verify the created group."
}}

Write-Output ("CREATED_DN=" + $group.DistinguishedName)
"""
    out = _run_ps(ps)

    created_dn = None
    for line in out.splitlines():
        if line.startswith("CREATED_DN="):
            created_dn = line.split("=", 1)[1].strip()
            break

    if not created_dn:
        raise ADServiceError(
            f"Group create executed but verification output not returned. STDOUT={out}"
        )

    conn = _connect(cfg)
    try:
        errors = []

        def _validate_dns(dns: Iterable[str], object_class: str, label: str) -> List[str]:
            valid = []
            for dn in dns or []:
                value = str(dn).strip()
                if not value:
                    continue
                if _entry_exists_by_dn(conn, value, object_class=object_class):
                    valid.append(value)
                elif not skip_invalid_members:
                    errors.append(f"{label} '{value}' not found.")
            seen = set()
            out_dns = []
            for dn in valid:
                if dn not in seen:
                    seen.add(dn)
                    out_dns.append(dn)
            return out_dns

        valid_user_dns = _validate_dns(user_dns, "user", "User")
        valid_computer_dns = _validate_dns(computer_dns, "computer", "Computer")
        valid_nested_group_dns = _validate_dns(nested_group_dns, "group", "Nested group")

        copied_member_dns: List[str] = []
        if copy_from_group_dn:
            if not _entry_exists_by_dn(conn, copy_from_group_dn, object_class="group"):
                raise ADServiceError(f"Copy source group not found: {copy_from_group_dn}")
            copied_member_dns = _get_group_members_dns(conn, copy_from_group_dn)

        if errors:
            raise ADServiceError(" ; ".join(errors))

        if owner_dn:
            if not _entry_exists_by_dn(conn, owner_dn):
                raise ADServiceError(f"Owner not found: {owner_dn}")
            _set_group_owner_ldap(conn, created_dn, owner_dn)

        all_member_dns: List[str] = []
        for seq in (valid_user_dns, valid_computer_dns, valid_nested_group_dns, copied_member_dns):
            for dn in seq:
                if dn not in all_member_dns:
                    all_member_dns.append(dn)

        _add_members_to_group_ldap(conn, created_dn, all_member_dns)
    finally:
        conn.unbind()

    if protect_from_deletion:
        _protect_group_from_deletion_winrm(cfg, created_dn)


def update_group(
    group_name: str,
    description: str = "",
    group_scope: str = "",
    group_category: str = "",
    add_member_dns: Optional[Iterable[str]] = None,
    remove_member_dns: Optional[Iterable[str]] = None,
) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return

    group_name = _safe_str(group_name)
    if not group_name:
        raise ADServiceError("Group name is required.")

    conn = _connect(cfg)
    try:
        group_dn = _get_group_dn(conn, cfg, group_name)
        if not group_dn:
            raise ADServiceError(f"Group not found: {group_name}")

        if _safe_str(description):
            ok = conn.modify(group_dn, {"description": [(MODIFY_REPLACE, [_safe_str(description)])]})
            if not ok:
                raise ADServiceError(f"Failed to update group description: {conn.result}")

        add_dns = [str(v).strip() for v in (add_member_dns or []) if str(v).strip()]
        if add_dns:
            _add_members_to_group_ldap(conn, group_dn, add_dns)

        remove_dns = [str(v).strip() for v in (remove_member_dns or []) if str(v).strip()]
        if remove_dns:
            _remove_members_from_group_ldap(conn, group_dn, remove_dns)
    finally:
        conn.unbind()

    if group_scope or group_category:
        _update_group_scope_category_winrm(
            cfg=cfg,
            group_dn=group_dn,
            group_scope=group_scope,
            group_category=group_category,
        )


def delete_group(group_name: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for group deletion.")

    group_name_esc = _ps_escape_single_quotes(group_name)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$group = Get-ADGroup -Identity '{group_name_esc}' -Server '{server_esc}' -ErrorAction Stop
Remove-ADGroup -Identity $group.DistinguishedName -Server '{server_esc}' -Confirm:$false
Write-Output "DELETED=1"
"""
    out = _run_ps(ps)
    if "DELETED=1" not in out:
        raise ADServiceError(f"Group delete executed but verification output not returned. STDOUT={out}")


def move_group(group_name: str, target_ou_dn: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for group move.")

    group_name = _safe_str(group_name)
    target_ou_dn = _safe_str(target_ou_dn)

    if not group_name:
        raise ADServiceError("Group name is required.")
    if not target_ou_dn:
        raise ADServiceError("Target OU DN is required.")

    conn = _connect(cfg)
    try:
        group_dn = _get_group_dn(conn, cfg, group_name)
        if not group_dn:
            raise ADServiceError(f"Group not found: {group_name}")

        ok = conn.search(
            search_base=group_dn,
            search_filter="(objectClass=group)",
            search_scope=BASE,
            attributes=["cn", "name", "distinguishedName"],
            size_limit=1,
        )
        if not ok or not conn.entries:
            raise ADServiceError(f"Could not read group before move: {group_name}")

        entry = conn.entries[0]
        cn = _entry_str(entry, "cn", "name") or group_name

        conflict_filter = f"(name={_escape_ldap_filter_value(cn)})"
        conn.search(
            search_base=target_ou_dn,
            search_filter=conflict_filter,
            search_scope=LEVEL,
            attributes=["distinguishedName", "objectClass", "name"],
            size_limit=10,
        )

        conflicts = [
            e.entry_dn
            for e in conn.entries
            if e.entry_dn.lower() != group_dn.lower()
        ]

        if conflicts:
            raise ADServiceError(
                f"An object named '{cn}' already exists in the target location: {conflicts[0]}"
            )
    finally:
        conn.unbind()

    group_dn_esc = _ps_escape_single_quotes(group_dn)
    target_ou_esc = _ps_escape_single_quotes(target_ou_dn)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

Move-ADObject -Identity '{group_dn_esc}' -TargetPath '{target_ou_esc}' -Server '{server_esc}'
Write-Output "MOVED=1"
"""
    out = _run_ps(ps)
    if "MOVED=1" not in out:
        raise ADServiceError(
            f"Group move executed but verification output not returned. STDOUT={out}"
        )


def add_group_members(group_name: str, members: Iterable[str]) -> None:
    update_group(group_name=group_name, add_member_dns=members)


def remove_group_members(group_name: str, members: Iterable[str]) -> None:
    update_group(group_name=group_name, remove_member_dns=members)


def add_user_to_groups(user_dn: str, group_dns: Iterable[str]) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return

    conn = _connect(cfg)
    try:
        user_dn = _safe_str(user_dn)
        if not user_dn:
            raise ADServiceError("User DN is required.")
        if not _entry_exists_by_dn(conn, user_dn, object_class="user"):
            raise ADServiceError(f"User not found: {user_dn}")

        for group_dn in [str(g).strip() for g in (group_dns or []) if str(g).strip()]:
            if not _entry_exists_by_dn(conn, group_dn, object_class="group"):
                continue
            _add_members_to_group_ldap(conn, group_dn, [user_dn])
    finally:
        conn.unbind()


# -----------------------------
# User operations via WinRM
# -----------------------------

def create_user(
    username: str,
    first_name: str,
    last_name: str,
    email: str,
    password: str,
    phone: str = "",
    department: str = "",
    description: str = "",
    hr_id: str = "",
    target_ou_dn: Optional[str] = None,
    group_dns: Optional[Iterable[str]] = None,
    must_change_password: bool = False,
    user_cannot_change_password: bool = False,
    password_never_expires: bool = False,
    account_disabled: bool = False,
) -> str:
    cfg = _load_config()
    if cfg.safe_mode:
        return ""
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for user creation.")

    username = _safe_str(username)
    first_name = _safe_str(first_name)
    last_name = _safe_str(last_name)
    email = _safe_str(email)
    password = _safe_str(password)
    phone = _safe_str(phone)
    department = _safe_str(department)
    description = _safe_str(description)
    hr_id = _safe_str(hr_id)
    ou_dn = _safe_str(target_ou_dn) or _safe_str(cfg.users_ou_dn)

    if not username:
        raise ADServiceError("Username is required.")
    if not password:
        raise ADServiceError("Password is required.")
    if not ou_dn:
        raise ADServiceError("Target OU DN is empty and no default Users OU DN is configured.")

    display_name = f"{first_name} {last_name}".strip() or username
    upn = _username_to_upn(username, cfg.upn_suffix)

    u_username = _ps_escape_single_quotes(username)
    u_first = _ps_escape_single_quotes(first_name or username)
    u_last = _ps_escape_single_quotes(last_name or username)
    u_display = _ps_escape_single_quotes(display_name)
    u_email = _ps_escape_single_quotes(email)
    u_password = _ps_escape_single_quotes(password)
    u_phone = _ps_escape_single_quotes(phone)
    u_department = _ps_escape_single_quotes(department)
    u_description = _ps_escape_single_quotes(description)
    u_hr_id = _ps_escape_single_quotes(hr_id)
    u_ou = _ps_escape_single_quotes(ou_dn)
    u_upn = _ps_escape_single_quotes(upn)
    u_server = _ps_escape_single_quotes(cfg.server_name)

    ps_must_change = "$true" if must_change_password else "$false"
    ps_password_never_expires = "$true" if password_never_expires else "$false"
    ps_enabled = "$false" if account_disabled else "$true"

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$server = '{u_server}'
$username = '{u_username}'
$ou = '{u_ou}'
$upn = '{u_upn}'

$existing = Get-ADUser -LDAPFilter "(sAMAccountName=$username)" -Server $server -ErrorAction SilentlyContinue
if ($existing) {{
    throw "User already exists: $($existing.DistinguishedName)"
}}

$sec = ConvertTo-SecureString '{u_password}' -AsPlainText -Force

$params = @{{
    Name                  = $username
    SamAccountName        = $username
    UserPrincipalName     = $upn
    GivenName             = '{u_first}'
    Surname               = '{u_last}'
    DisplayName           = '{u_display}'
    Path                  = $ou
    AccountPassword       = $sec
    Server                = $server
    Enabled               = {ps_enabled}
    ChangePasswordAtLogon = {ps_must_change}
    PasswordNeverExpires  = {ps_password_never_expires}
}}

if ('{u_email}')       {{ $params['EmailAddress'] = '{u_email}' }}
if ('{u_description}') {{ $params['Description'] = '{u_description}' }}
if ('{u_department}')  {{ $params['Department'] = '{u_department}' }}
if ('{u_phone}')       {{ $params['OfficePhone'] = '{u_phone}' }}
if ('{u_hr_id}')       {{ $params['EmployeeID'] = '{u_hr_id}' }}

New-ADUser @params

$created = Get-ADUser -Identity '{u_username}' -Server $server -Properties DistinguishedName
if (-not $created) {{
    throw "User creation succeeded but could not verify the created user."
}}

Write-Output ("CREATED_DN=" + $created.DistinguishedName)
"""
    out = _run_ps(ps)

    created_dn = None
    for line in out.splitlines():
        if line.startswith("CREATED_DN="):
            created_dn = line.split("=", 1)[1].strip()
            break

    if not created_dn:
        raise ADServiceError(f"User was created but DN was not returned. STDOUT={out}")

    if group_dns:
        add_user_to_groups(created_dn, group_dns)

    return created_dn


def update_user(username: str, updates: dict) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for user update.")

    username_esc = _ps_escape_single_quotes(_safe_str(username))
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    if not username_esc:
        raise ADServiceError("Username is required.")

    clean_updates = {}
    for key, value in (updates or {}).items():
        if value is None or value == "":
            continue
        clean_updates[key] = _ps_escape_single_quotes(value)

    if not clean_updates:
        return

    replace_lines = []
    for key, value in clean_updates.items():
        replace_lines.append(f"$replace['{key}'] = '{value}'")
    replace_block = "\n".join(replace_lines)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$user = Get-ADUser -Identity '{username_esc}' -Server '{server_esc}' -ErrorAction Stop
$replace = @{{}}
{replace_block}

Set-ADUser -Identity $user.DistinguishedName -Server '{server_esc}' -Replace $replace
Write-Output ("UPDATED_DN=" + $user.DistinguishedName)
"""
    out = _run_ps(ps)
    if "UPDATED_DN=" not in out:
        raise ADServiceError(f"Update executed but verification output not returned. STDOUT={out}")


def reset_user_password(username: str, new_password: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for password reset.")

    username_esc = _ps_escape_single_quotes(username)
    password_esc = _ps_escape_single_quotes(new_password)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$sec = ConvertTo-SecureString '{password_esc}' -AsPlainText -Force
$user = Get-ADUser -Identity '{username_esc}' -Server '{server_esc}' -ErrorAction Stop

Set-ADAccountPassword -Identity $user.DistinguishedName -Server '{server_esc}' -Reset -NewPassword $sec
Set-ADUser -Identity $user.DistinguishedName -Server '{server_esc}' -ChangePasswordAtLogon $true
Write-Output "PASSWORD_RESET=1"
"""
    out = _run_ps(ps)
    if "PASSWORD_RESET=1" not in out:
        raise ADServiceError(f"Password reset executed but no verification output returned. STDOUT={out}")


def lock_user(username: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for user lock.")

    username_esc = _ps_escape_single_quotes(_safe_str(username))
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$user = Get-ADUser -Identity '{username_esc}' -Server '{server_esc}' -ErrorAction Stop
Disable-ADAccount -Identity $user.DistinguishedName -Server '{server_esc}'
Write-Output "LOCKED=1"
"""
    out = _run_ps(ps)
    if "LOCKED=1" not in out:
        raise ADServiceError(f"Account lock executed but verification output not returned. STDOUT={out}")


def unlock_user(username: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for user unlock.")

    username_esc = _ps_escape_single_quotes(_safe_str(username))
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$user = Get-ADUser -Identity '{username_esc}' -Server '{server_esc}' -ErrorAction Stop
Enable-ADAccount -Identity $user.DistinguishedName -Server '{server_esc}'
Write-Output "UNLOCKED=1"
"""
    out = _run_ps(ps)
    if "UNLOCKED=1" not in out:
        raise ADServiceError(f"Account unlock executed but verification output not returned. STDOUT={out}")


def move_user(username: str, target_ou_dn: str) -> str:
    cfg = _load_config()
    if cfg.safe_mode:
        return ""
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for user move.")

    valid_ou_dns = {ou["dn"] for ou in list_ous() if ou.get("dn")}
    if target_ou_dn not in valid_ou_dns:
        raise ADServiceError("Selected OU is invalid.")

    username_esc = _ps_escape_single_quotes(_safe_str(username))
    target_ou_esc = _ps_escape_single_quotes(_safe_str(target_ou_dn))
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$user = Get-ADUser -Identity '{username_esc}' -Server '{server_esc}' -Properties DistinguishedName -ErrorAction Stop
Move-ADObject -Identity $user.DistinguishedName -TargetPath '{target_ou_esc}' -Server '{server_esc}'
$verify = Get-ADUser -Identity '{username_esc}' -Server '{server_esc}' -Properties DistinguishedName
Write-Output ("MOVED_DN=" + $verify.DistinguishedName)
"""
    out = _run_ps(ps)

    moved_dn = None
    for line in out.splitlines():
        if line.startswith("MOVED_DN="):
            moved_dn = line.split("=", 1)[1].strip()
            break

    if not moved_dn:
        raise ADServiceError(f"Move executed but verification output not returned. STDOUT={out}")

    return moved_dn


# -----------------------------
# Computer operations via WinRM
# -----------------------------

def create_computer(computer_name: str, ou_dn: Optional[str] = None, description: str = "") -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for computer creation.")

    name = _safe_str(computer_name)
    if not name:
        raise ADServiceError("Computer name is required.")

    target_ou = _safe_str(ou_dn) or _safe_str(cfg.computers_ou_dn)
    if not target_ou:
        raise ADServiceError("Target OU DN is required.")

    name_esc = _ps_escape_single_quotes(name)
    ou_esc = _ps_escape_single_quotes(target_ou)
    desc_esc = _ps_escape_single_quotes(description)
    sam_esc = _ps_escape_single_quotes(_normalize_computer_sam(name))
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$existing = Get-ADComputer -Identity '{name_esc}' -Server '{server_esc}' -ErrorAction SilentlyContinue
if ($existing) {{
    throw "Computer already exists: $($existing.DistinguishedName)"
}}

$params = @{{
    Name           = '{name_esc}'
    SamAccountName = '{sam_esc}'
    Path           = '{ou_esc}'
    Server         = '{server_esc}'
}}

if ('{desc_esc}') {{ $params['Description'] = '{desc_esc}' }}

New-ADComputer @params
$verify = Get-ADComputer -Identity '{name_esc}' -Server '{server_esc}' -Properties DistinguishedName
Write-Output ("CREATED_DN=" + $verify.DistinguishedName)
"""
    out = _run_ps(ps)
    if "CREATED_DN=" not in out:
        raise ADServiceError(f"Computer create executed but verification output not returned. STDOUT={out}")


def lock_computer(computer_name: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for computer lock.")

    name_esc = _ps_escape_single_quotes(computer_name)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$computer = Get-ADComputer -Identity '{name_esc}' -Server '{server_esc}' -ErrorAction Stop
Disable-ADAccount -Identity $computer.DistinguishedName -Server '{server_esc}'
Write-Output "LOCKED=1"
"""
    out = _run_ps(ps)
    if "LOCKED=1" not in out:
        raise ADServiceError(f"Computer lock executed but verification output not returned. STDOUT={out}")


def unlock_computer(computer_name: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for computer unlock.")

    name_esc = _ps_escape_single_quotes(computer_name)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$computer = Get-ADComputer -Identity '{name_esc}' -Server '{server_esc}' -ErrorAction Stop
Enable-ADAccount -Identity $computer.DistinguishedName -Server '{server_esc}'
Write-Output "UNLOCKED=1"
"""
    out = _run_ps(ps)
    if "UNLOCKED=1" not in out:
        raise ADServiceError(f"Computer unlock executed but verification output not returned. STDOUT={out}")


def move_computer(computer_name: str, target_ou_dn: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for computer move.")

    name_esc = _ps_escape_single_quotes(computer_name)
    target_ou_esc = _ps_escape_single_quotes(target_ou_dn)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$computer = Get-ADComputer -Identity '{name_esc}' -Server '{server_esc}' -Properties DistinguishedName -ErrorAction Stop
Move-ADObject -Identity $computer.DistinguishedName -TargetPath '{target_ou_esc}' -Server '{server_esc}'
Write-Output "MOVED=1"
"""
    out = _run_ps(ps)
    if "MOVED=1" not in out:
        raise ADServiceError(f"Computer move executed but verification output not returned. STDOUT={out}")


def update_computer(computer_name: str, updates: dict) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for computer update.")

    name_esc = _ps_escape_single_quotes(computer_name)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    clean_updates = {}
    for key, value in (updates or {}).items():
        if value is None or value == "":
            continue
        clean_updates[key] = _ps_escape_single_quotes(value)

    if not clean_updates:
        return

    replace_lines = []
    for key, value in clean_updates.items():
        replace_lines.append(f"$replace['{key}'] = '{value}'")
    replace_block = "\n".join(replace_lines)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$computer = Get-ADComputer -Identity '{name_esc}' -Server '{server_esc}' -ErrorAction Stop
$replace = @{{}}
{replace_block}

Set-ADComputer -Identity $computer.DistinguishedName -Server '{server_esc}' -Replace $replace
Write-Output "UPDATED=1"
"""
    out = _run_ps(ps)
    if "UPDATED=1" not in out:
        raise ADServiceError(f"Computer update executed but verification output not returned. STDOUT={out}")


# -----------------------------
# OU operations via WinRM
# -----------------------------

def create_ou(ou_name: str, parent_dn: str, description: str = "", protect: bool = False) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for OU creation.")

    ou_name_esc = _ps_escape_single_quotes(ou_name)
    parent_base = _safe_str(parent_dn) or _safe_str(cfg.base_dn)
    if not parent_base:
        raise ADServiceError("Parent OU DN or base DN is required for OU creation.")

    parent_esc = _ps_escape_single_quotes(parent_base)
    description_esc = _ps_escape_single_quotes(description)
    server_esc = _ps_escape_single_quotes(cfg.server_name)
    ps_protect = "$true" if protect else "$false"

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$params = @{{
    Name   = '{ou_name_esc}'
    Path   = '{parent_esc}'
    Server = '{server_esc}'
}}

if ('{description_esc}') {{ $params['Description'] = '{description_esc}' }}
$params['ProtectedFromAccidentalDeletion'] = {ps_protect}

New-ADOrganizationalUnit @params
Write-Output "CREATED=1"
"""
    out = _run_ps(ps)
    if "CREATED=1" not in out:
        raise ADServiceError(f"OU create executed but verification output not returned. STDOUT={out}")


def update_ou(ou_dn: str, new_name: str = "", description: str = "", protect: bool = False) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for OU update.")

    ou_dn_esc = _ps_escape_single_quotes(ou_dn)
    new_name_esc = _ps_escape_single_quotes(new_name)
    description_esc = _ps_escape_single_quotes(description)
    server_esc = _ps_escape_single_quotes(cfg.server_name)
    ps_protect = "$true" if protect else "$false"

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

$targetDn = '{ou_dn_esc}'

if ('{new_name_esc}') {{
    Rename-ADObject -Identity $targetDn -NewName '{new_name_esc}' -Server '{server_esc}'
    $obj = Get-ADObject -Identity $targetDn -Server '{server_esc}' -Properties DistinguishedName
    $parentDn = $obj.DistinguishedName.Split(',', 2)[1]
    $targetDn = "OU={new_name_esc}," + $parentDn
}}

if ('{description_esc}') {{
    Set-ADOrganizationalUnit -Identity $targetDn -Server '{server_esc}' -Description '{description_esc}'
}}

Set-ADOrganizationalUnit -Identity $targetDn -Server '{server_esc}' -ProtectedFromAccidentalDeletion {ps_protect}
Write-Output ("UPDATED_DN=" + $targetDn)
"""
    out = _run_ps(ps)
    if "UPDATED_DN=" not in out:
        raise ADServiceError(f"OU update executed but verification output not returned. STDOUT={out}")


def delete_ou(ou_dn: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for OU deletion.")

    ou_dn_esc = _ps_escape_single_quotes(ou_dn)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

Set-ADOrganizationalUnit -Identity '{ou_dn_esc}' -Server '{server_esc}' -ProtectedFromAccidentalDeletion $false -ErrorAction SilentlyContinue
Remove-ADOrganizationalUnit -Identity '{ou_dn_esc}' -Server '{server_esc}' -Recursive -Confirm:$false
Write-Output "DELETED=1"
"""
    out = _run_ps(ps)
    if "DELETED=1" not in out:
        raise ADServiceError(f"OU delete executed but verification output not returned. STDOUT={out}")


def move_ou(ou_dn: str, target_parent_dn: str) -> None:
    cfg = _load_config()
    if cfg.safe_mode:
        return
    if not cfg.server_name:
        raise ADServiceError("WinRM server_name is required for OU move.")

    ou_dn_esc = _ps_escape_single_quotes(ou_dn)
    target_parent_esc = _ps_escape_single_quotes(target_parent_dn)
    server_esc = _ps_escape_single_quotes(cfg.server_name)

    ps = f"""
$ErrorActionPreference = 'Stop'
Import-Module ActiveDirectory

Move-ADObject -Identity '{ou_dn_esc}' -TargetPath '{target_parent_esc}' -Server '{server_esc}'
Write-Output "MOVED=1"
"""
    out = _run_ps(ps)
    if "MOVED=1" not in out:
        raise ADServiceError(f"OU move executed but verification output not returned. STDOUT={out}")


# -----------------------------
# Bulk group creation from Excel
# -----------------------------

def _split_excel_multi_value(value) -> list[str]:
    if value is None:
        return []
    return [item.strip() for item in str(value).split(";") if item and str(item).strip()]


def _to_bool_excel(value, default=False) -> bool:
    if value is None or str(value).strip() == "":
        return default
    return str(value).strip().lower() in {"1", "true", "yes", "y"}


def parse_bulk_group_excel(file_obj) -> list[dict]:
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    normalized_headers = [h.lower() for h in headers]

    if "group_name" not in normalized_headers:
        raise ADServiceError("Missing required Excel column: group_name")

    parsed = []
    for idx, row in enumerate(rows[1:], start=2):
        row_dict = dict(zip(normalized_headers, row))
        if not any(v is not None and str(v).strip() != "" for v in row_dict.values()):
            continue

        parsed.append(
            {
                "excel_row": idx,
                "group_name": str(row_dict.get("group_name") or "").strip(),
                "description": str(row_dict.get("description") or "").strip(),
                "ou_name": str(row_dict.get("ou_name") or "").strip(),
                "group_scope": str(row_dict.get("group_scope") or "Global").strip() or "Global",
                "group_category": str(row_dict.get("group_category") or "Security").strip() or "Security",
                "owner_username": str(row_dict.get("owner_username") or "").strip(),
                "users": _split_excel_multi_value(row_dict.get("users")),
                "computers": _split_excel_multi_value(row_dict.get("computers")),
                "nested_groups": _split_excel_multi_value(row_dict.get("nested_groups")),
                "copy_from_group": str(row_dict.get("copy_from_group") or "").strip(),
                "protect_from_deletion": _to_bool_excel(row_dict.get("protect_from_deletion"), default=False),
            }
        )

    return parsed


def _resolve_bulk_group_row(conn: Connection, cfg: ADConfig, row: dict) -> dict:
    errors: list[str] = []

    group_name = _safe_str(row.get("group_name"))
    if not group_name:
        errors.append("group_name is required")

    group_scope = _safe_str(row.get("group_scope") or "Global") or "Global"
    if group_scope not in {"Global", "DomainLocal", "Universal"}:
        errors.append(f"Invalid group_scope '{group_scope}'")

    group_category = _safe_str(row.get("group_category") or "Security") or "Security"
    if group_category not in {"Security", "Distribution"}:
        errors.append(f"Invalid group_category '{group_category}'")

    ou_dn = cfg.groups_ou_dn
    ou_name = _safe_str(row.get("ou_name"))
    if ou_name:
        try:
            ou_dn = resolve_ou_name_to_dn(ou_name)
        except ADServiceError as e:
            errors.append(str(e))

    owner_dn = ""
    owner_username = _safe_str(row.get("owner_username"))
    if owner_username:
        try:
            owner_dn = resolve_user_username_to_dn(conn, cfg, owner_username)
        except ADServiceError as e:
            errors.append(str(e))

    user_dns = []
    for username in row.get("users", []):
        try:
            user_dns.append(resolve_user_username_to_dn(conn, cfg, username))
        except ADServiceError:
            errors.append(f"User '{username}' not found.")

    computer_dns = []
    for computer_name in row.get("computers", []):
        try:
            computer_dns.append(resolve_computer_name_to_dn(conn, cfg, computer_name))
        except ADServiceError:
            errors.append(f"Computer '{computer_name}' not found.")

    nested_group_dns = []
    for nested_group_name in row.get("nested_groups", []):
        try:
            nested_group_dns.append(resolve_group_name_to_dn(conn, cfg, nested_group_name))
        except ADServiceError:
            errors.append(f"Nested group '{nested_group_name}' not found.")

    copy_from_group_dn = ""
    copy_from_group = _safe_str(row.get("copy_from_group"))
    if copy_from_group:
        try:
            copy_from_group_dn = resolve_group_name_to_dn(conn, cfg, copy_from_group)
        except ADServiceError:
            errors.append(f"Copy source group '{copy_from_group}' not found.")

    return {
        "excel_row": row.get("excel_row"),
        "group_name": group_name,
        "description": _safe_str(row.get("description")),
        "group_scope": group_scope,
        "group_category": group_category,
        "ou_dn": ou_dn,
        "owner_dn": owner_dn,
        "user_dns": user_dns,
        "computer_dns": computer_dns,
        "nested_group_dns": nested_group_dns,
        "copy_from_group_dn": copy_from_group_dn,
        "protect_from_deletion": bool(row.get("protect_from_deletion")),
        "errors": errors,
    }


def bulk_create_groups_from_excel(file_obj) -> list[dict]:
    cfg = _load_config()
    rows = parse_bulk_group_excel(file_obj)
    results = []

    if cfg.safe_mode:
        for row in rows:
            results.append(
                {
                    "excel_row": row.get("excel_row"),
                    "group_name": row.get("group_name", ""),
                    "success": True,
                    "message": "Safe mode enabled, no changes were applied.",
                }
            )
        return results

    conn = _connect(cfg)
    try:
        for row in rows:
            resolved = _resolve_bulk_group_row(conn, cfg, row)
            if resolved["errors"]:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": " ; ".join(resolved["errors"]),
                    }
                )
                continue

            try:
                create_group(
                    group_name=resolved["group_name"],
                    description=resolved["description"],
                    user_dns=resolved["user_dns"],
                    computer_dns=resolved["computer_dns"],
                    nested_group_dns=resolved["nested_group_dns"],
                    group_scope=resolved["group_scope"],
                    group_category=resolved["group_category"],
                    owner_dn=resolved["owner_dn"],
                    ou_dn=resolved["ou_dn"],
                    protect_from_deletion=resolved["protect_from_deletion"],
                    copy_from_group_dn=resolved["copy_from_group_dn"],
                    skip_invalid_members=False,
                )
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": True,
                        "message": "Created successfully.",
                    }
                )
            except Exception as e:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": str(e),
                    }
                )
    finally:
        conn.unbind()

    return results


# -----------------------------
# Bulk group update from Excel
# -----------------------------

def parse_bulk_group_update_excel(file_obj) -> list[dict]:
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    normalized_headers = [h.lower() for h in headers]

    if "group_name" not in normalized_headers:
        raise ADServiceError("Missing required Excel column: group_name")

    parsed = []
    for idx, row in enumerate(rows[1:], start=2):
        row_dict = dict(zip(normalized_headers, row))
        if not any(v is not None and str(v).strip() != "" for v in row_dict.values()):
            continue

        parsed.append(
            {
                "excel_row": idx,
                "group_name": str(row_dict.get("group_name") or "").strip(),
                "description": str(row_dict.get("description") or "").strip(),
                "group_scope": str(row_dict.get("group_scope") or "").strip(),
                "group_category": str(row_dict.get("group_category") or "").strip(),
                "add_users": _split_excel_multi_value(row_dict.get("add_users")),
                "add_computers": _split_excel_multi_value(row_dict.get("add_computers")),
                "add_nested_groups": _split_excel_multi_value(row_dict.get("add_nested_groups")),
                "remove_users": _split_excel_multi_value(row_dict.get("remove_users")),
                "remove_computers": _split_excel_multi_value(row_dict.get("remove_computers")),
                "remove_nested_groups": _split_excel_multi_value(row_dict.get("remove_nested_groups")),
            }
        )

    return parsed


def _resolve_bulk_group_update_row(conn: Connection, cfg: ADConfig, row: dict) -> dict:
    errors: list[str] = []

    group_name = _safe_str(row.get("group_name"))
    if not group_name:
        errors.append("group_name is required")

    group_dn = ""
    if group_name:
        try:
            group_dn = resolve_group_name_to_dn(conn, cfg, group_name)
        except ADServiceError as e:
            errors.append(str(e))

    description = _safe_str(row.get("description"))

    raw_scope = _safe_str(row.get("group_scope"))
    raw_category = _safe_str(row.get("group_category"))

    normalized_scope = ""
    normalized_category = ""

    if raw_scope:
        try:
            normalized_scope = _normalize_group_scope(raw_scope)
        except ADServiceError:
            errors.append(f"Invalid group_scope '{raw_scope}'")

    if raw_category:
        try:
            normalized_category = _normalize_group_category(raw_category)
        except ADServiceError:
            errors.append(f"Invalid group_category '{raw_category}'")

    add_member_dns: list[str] = []
    remove_member_dns: list[str] = []

    for username in row.get("add_users", []):
        try:
            add_member_dns.append(resolve_user_username_to_dn(conn, cfg, username))
        except ADServiceError:
            errors.append(f"Add user '{username}' not found.")

    for computer_name in row.get("add_computers", []):
        try:
            add_member_dns.append(resolve_computer_name_to_dn(conn, cfg, computer_name))
        except ADServiceError:
            errors.append(f"Add computer '{computer_name}' not found.")

    for nested_group_name in row.get("add_nested_groups", []):
        try:
            add_member_dns.append(resolve_group_name_to_dn(conn, cfg, nested_group_name))
        except ADServiceError:
            errors.append(f"Add nested group '{nested_group_name}' not found.")

    for username in row.get("remove_users", []):
        try:
            remove_member_dns.append(resolve_user_username_to_dn(conn, cfg, username))
        except ADServiceError:
            errors.append(f"Remove user '{username}' not found.")

    for computer_name in row.get("remove_computers", []):
        try:
            remove_member_dns.append(resolve_computer_name_to_dn(conn, cfg, computer_name))
        except ADServiceError:
            errors.append(f"Remove computer '{computer_name}' not found.")

    for nested_group_name in row.get("remove_nested_groups", []):
        try:
            remove_member_dns.append(resolve_group_name_to_dn(conn, cfg, nested_group_name))
        except ADServiceError:
            errors.append(f"Remove nested group '{nested_group_name}' not found.")

    return {
        "excel_row": row.get("excel_row"),
        "group_name": group_name,
        "group_dn": group_dn,
        "description": description,
        "group_scope": normalized_scope,
        "group_category": normalized_category,
        "add_member_dns": add_member_dns,
        "remove_member_dns": remove_member_dns,
        "errors": errors,
    }


def bulk_update_groups_from_excel(file_obj) -> list[dict]:
    cfg = _load_config()
    rows = parse_bulk_group_update_excel(file_obj)
    results = []

    if cfg.safe_mode:
        for row in rows:
            results.append(
                {
                    "excel_row": row.get("excel_row"),
                    "group_name": row.get("group_name", ""),
                    "success": True,
                    "message": "Safe mode enabled, no changes were applied.",
                }
            )
        return results

    conn = _connect(cfg)
    try:
        for row in rows:
            resolved = _resolve_bulk_group_update_row(conn, cfg, row)

            if resolved["errors"]:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": " ; ".join(resolved["errors"]),
                    }
                )
                continue

            try:
                update_group(
                    group_name=resolved["group_name"],
                    description=resolved["description"],
                    group_scope=resolved["group_scope"],
                    group_category=resolved["group_category"],
                    add_member_dns=resolved["add_member_dns"],
                    remove_member_dns=resolved["remove_member_dns"],
                )
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": True,
                        "message": "Updated successfully.",
                    }
                )
            except Exception as e:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": str(e),
                    }
                )
    finally:
        conn.unbind()

    return results


# -----------------------------
# Bulk group move/delete from Excel
# -----------------------------

def parse_bulk_group_move_excel(file_obj) -> list[dict]:
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    normalized_headers = [h.lower() for h in headers]

    required = {"group_name", "target_ou_name"}
    missing = [col for col in required if col not in normalized_headers]
    if missing:
        raise ADServiceError(f"Missing required Excel columns: {', '.join(missing)}")

    parsed = []
    for idx, row in enumerate(rows[1:], start=2):
        row_dict = dict(zip(normalized_headers, row))
        if not any(v is not None and str(v).strip() != "" for v in row_dict.values()):
            continue

        parsed.append(
            {
                "excel_row": idx,
                "group_name": str(row_dict.get("group_name") or "").strip(),
                "target_ou_name": str(row_dict.get("target_ou_name") or "").strip(),
            }
        )

    return parsed


def parse_bulk_group_delete_excel(file_obj) -> list[dict]:
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    normalized_headers = [h.lower() for h in headers]

    if "group_name" not in normalized_headers:
        raise ADServiceError("Missing required Excel column: group_name")

    parsed = []
    for idx, row in enumerate(rows[1:], start=2):
        row_dict = dict(zip(normalized_headers, row))
        if not any(v is not None and str(v).strip() != "" for v in row_dict.values()):
            continue

        parsed.append(
            {
                "excel_row": idx,
                "group_name": str(row_dict.get("group_name") or "").strip(),
            }
        )

    return parsed


def _resolve_bulk_group_move_row(conn: Connection, cfg: ADConfig, row: dict) -> dict:
    errors: list[str] = []

    group_name = _safe_str(row.get("group_name"))
    target_ou_name = _safe_str(row.get("target_ou_name"))

    if not group_name:
        errors.append("group_name is required")

    if not target_ou_name:
        errors.append("target_ou_name is required")

    group_dn = ""
    if group_name:
        dn = _get_group_dn(conn, cfg, group_name)
        if not dn:
            errors.append(f"Group '{group_name}' not found.")
        else:
            group_dn = dn

    target_ou_dn = ""
    if target_ou_name:
        try:
            target_ou_dn = resolve_ou_name_to_dn(target_ou_name)
        except ADServiceError as e:
            errors.append(str(e))

    return {
        "excel_row": row.get("excel_row"),
        "group_name": group_name,
        "group_dn": group_dn,
        "target_ou_name": target_ou_name,
        "target_ou_dn": target_ou_dn,
        "errors": errors,
    }


def _resolve_bulk_group_delete_row(conn: Connection, cfg: ADConfig, row: dict) -> dict:
    errors: list[str] = []

    group_name = _safe_str(row.get("group_name"))
    if not group_name:
        errors.append("group_name is required")

    group_dn = ""
    if group_name:
        dn = _get_group_dn(conn, cfg, group_name)
        if not dn:
            errors.append(f"Group '{group_name}' not found.")
        else:
            group_dn = dn

    return {
        "excel_row": row.get("excel_row"),
        "group_name": group_name,
        "group_dn": group_dn,
        "errors": errors,
    }


def bulk_move_groups_from_excel(file_obj) -> list[dict]:
    cfg = _load_config()
    rows = parse_bulk_group_move_excel(file_obj)
    results = []

    if cfg.safe_mode:
        for row in rows:
            results.append(
                {
                    "excel_row": row.get("excel_row"),
                    "group_name": row.get("group_name", ""),
                    "success": True,
                    "message": "Safe mode enabled, no changes were applied.",
                }
            )
        return results

    conn = _connect(cfg)
    try:
        for row in rows:
            resolved = _resolve_bulk_group_move_row(conn, cfg, row)

            if resolved["errors"]:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": " ; ".join(resolved["errors"]),
                    }
                )
                continue

            try:
                move_group(
                    group_name=resolved["group_name"],
                    target_ou_dn=resolved["target_ou_dn"],
                )
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": True,
                        "message": f"Moved successfully to '{resolved['target_ou_name']}'.",
                    }
                )
            except Exception as e:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": str(e),
                    }
                )
    finally:
        conn.unbind()

    return results


def bulk_delete_groups_from_excel(file_obj) -> list[dict]:
    cfg = _load_config()
    rows = parse_bulk_group_delete_excel(file_obj)
    results = []

    if cfg.safe_mode:
        for row in rows:
            results.append(
                {
                    "excel_row": row.get("excel_row"),
                    "group_name": row.get("group_name", ""),
                    "success": True,
                    "message": "Safe mode enabled, no changes were applied.",
                }
            )
        return results

    conn = _connect(cfg)
    try:
        for row in rows:
            resolved = _resolve_bulk_group_delete_row(conn, cfg, row)

            if resolved["errors"]:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": " ; ".join(resolved["errors"]),
                    }
                )
                continue

            try:
                delete_group(resolved["group_name"])
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": True,
                        "message": "Deleted successfully.",
                    }
                )
            except Exception as e:
                results.append(
                    {
                        "excel_row": resolved["excel_row"],
                        "group_name": resolved["group_name"],
                        "success": False,
                        "message": str(e),
                    }
                )
    finally:
        conn.unbind()

    return results