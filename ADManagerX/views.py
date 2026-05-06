from django.conf import settings
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.http import FileResponse, Http404, JsonResponse
from django.shortcuts import render as django_render, redirect
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_http_methods, require_POST
import json
from functools import wraps
from ldap3 import Server, Connection, Tls, BASE, ALL
import os
import ssl

from .forms import LdapSettingsForm
from .models import LdapSettings, HelpdeskProfile
from .ad_service import (
    ADServiceError,
    bulk_create_groups_from_excel,
    bulk_update_groups_from_excel,
    bulk_move_groups_from_excel,
    bulk_delete_groups_from_excel,
    create_user,
    test_ldap_health,
    test_winrm_health,
    set_current_ldap_settings,
    clear_current_ldap_settings,
    reset_user_password,
    lock_user,
    unlock_user,
    move_user,
    update_user,
    create_computer,
    lock_computer,
    unlock_computer,
    move_computer,
    update_computer,
    create_group,
    update_group,
    delete_group,
    move_group,
    create_ou,
    update_ou,
    delete_ou as ad_delete_ou,
    move_ou as ad_move_ou,
    list_ous,
    list_groups,
    list_group_ous,
    resolve_ou_name_to_dn,
    search_computers,
    search_users,
    search_groups,
    search_directory_objects,
    get_group_members_for_ui,
    get_computer_report_data,
    get_ou_report_data,
)


def _read_excel_rows(uploaded_file):
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise ADServiceError("openpyxl is required for Excel uploads.") from exc

    if uploaded_file is None:
        raise ADServiceError("Please upload an Excel file.")

    wb = load_workbook(uploaded_file, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    data = []

    for row in rows[1:]:
        if not any(row):
            continue

        item = {}
        for idx, value in enumerate(row):
            key = headers[idx] if idx < len(headers) else f"col_{idx}"
            item[key] = value.strip() if isinstance(value, str) else value
        data.append(item)

    return data



def _validate_required_headers(rows, required_headers):
    if not rows:
        raise ADServiceError("No data found in Excel file.")

    first_row_keys = set(rows[0].keys())
    missing = [h for h in required_headers if h not in first_row_keys]
    if missing:
        raise ADServiceError(f"Missing required columns: {', '.join(missing)}")



def _get_uploaded_file(request, *keys):
    for key in keys:
        file_obj = request.FILES.get(key)
        if file_obj:
            return file_obj
    return None


# -----------------------------
# Authentication / authorization helpers
# -----------------------------

PERMISSION_DENIED_MESSAGE = "You do not have permission to perform this action."


def _get_user_profile(user):
    """
    Return the HelpdeskProfile for the logged-in Django user.

    This intentionally uses a direct query instead of user.helpdeskprofile because
    custom user models / related_name changes can make the reverse attribute fail.
    """
    if not user or not user.is_authenticated:
        return None

    try:
        return HelpdeskProfile.objects.filter(user=user).first()
    except Exception:
        return None


def _user_is_system_admin(user) -> bool:
    if not user or not user.is_authenticated:
        return False

    if user.is_superuser or user.is_staff:
        return True

    profile = _get_user_profile(user)
    return bool(profile and (profile.role or "").lower() == "admin")


def user_has_ad_permission(user, resource: str, action: str, domain_id=None) -> bool:
    """
    Check ADManagerX saved permissions.

    Superusers, staff users, and HelpdeskProfile role='admin' can access everything.
    The role defaults are also honored as a safety fallback, so an auditor can
    access report pages even if an older profile has empty permissions saved.
    """
    if not user or not user.is_authenticated:
        return False

    if _user_is_system_admin(user):
        return True

    profile = _get_user_profile(user)
    if not profile:
        return False

    role = _domain_role_for_user(user, domain_id) or "custom"

    permissions = _get_domain_permissions(profile, domain_id)
    allowed_actions = permissions.get(resource, []) or []

    role_defaults = _default_permissions_for_role(role)
    default_actions = role_defaults.get(resource, []) or []

    if action == "any":
        return bool(allowed_actions or default_actions)

    return action in allowed_actions or action in default_actions


def _first_post_value(request, keys: list[str]) -> str:
    for key in keys:
        value = (request.POST.get(key) or "").strip()
        if value:
            return value
    return ""


def _scope_allows_dn(user, target_dn: str, domain_id=None) -> bool:
    """
    Optional OU scope check.

    If profile.scope is empty, the user is not scope-limited.
    If target_dn is empty, we cannot validate scope here, so we allow the view
    and rely on the resource/action permission.
    If both exist, target_dn must be inside the profile.scope DN.
    """
    if not target_dn:
        return True

    if _user_is_system_admin(user):
        return True

    profile = _get_user_profile(user)
    if not profile:
        return False

    scope = _get_domain_scope(profile, domain_id)
    if not scope:
        return True

    return target_dn.lower().endswith(scope.lower())


def _default_redirect_for_resource(resource: str) -> str:
    """
    Always redirect permission failures to the dashboard.

    Do not redirect back to resource landing pages because those pages may also
    be protected, which can cause redirect loops or blank-looking pages.
    """
    return "home"


def ad_permission_required(resource: str, action: str = "any", scope_keys: list[str] | None = None):
    """
    Use this on ADManagerX views to enforce HelpdeskProfile.permissions.

    Example:
        @ad_permission_required("users", "create")
        def create_single_user(request):
            ...

    scope_keys are POST field names that may contain a target DN.
    When a helpdesk user has profile.scope set, those DN values must be inside
    that scope. This is mainly useful for create/move operations where the
    target OU DN is posted by the form.
    """
    scope_keys = scope_keys or [
        "target_ou_dn",
        "target_ou",
        "ou_dn",
        "ou",
        "target_parent_dn",
        "parent_ou",
        "scope",
    ]

    def decorator(view_func):
        @wraps(view_func)
        def wrapper(request, *args, **kwargs):
            if not request.user.is_authenticated:
                return redirect(f"{settings.LOGIN_URL}?next={request.get_full_path()}")

            active_domain = get_active_ldap_settings(request)
            domain_id = active_domain.id if active_domain else None

            if not user_has_ad_permission(request.user, resource, action, domain_id=domain_id):
                messages.error(request, f"You do not have permission to {action.replace('_', ' ')} {resource}.")
                return redirect(_default_redirect_for_resource(resource))

            if request.method == "POST":
                target_dn = _first_post_value(request, scope_keys)
                if target_dn and not _scope_allows_dn(request.user, target_dn, domain_id=domain_id):
                    messages.error(request, "This object is outside your allowed OU scope.")
                    return redirect(_default_redirect_for_resource(resource))

            return view_func(request, *args, **kwargs)

        return wrapper

    return decorator


def system_admin_required(view_func):
    """Allow only superuser/staff/profile role admin to access system administration."""
    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect(f"{settings.LOGIN_URL}?next={request.get_full_path()}")

        get_active_ldap_settings(request)

        if not _user_is_system_admin(request.user):
            messages.error(request, "Only system administrators can access this page.")
            return redirect("home")

        return view_func(request, *args, **kwargs)

    return wrapper




def _profile_domain_data(profile, domain_id):
    if not profile or not domain_id:
        return {}
    domain_permissions = profile.domain_permissions or {}
    return domain_permissions.get(str(domain_id), {}) or {}


def _domain_role_for_user(user, domain_id=None):
    profile = _get_user_profile(user)
    if not profile:
        return ""
    if domain_id:
        data = _profile_domain_data(profile, domain_id)
        if data.get("role"):
            return (data.get("role") or "").strip().lower()
    return (profile.role or "custom").strip().lower()


def _get_domain_permissions(profile, domain_id=None):
    if not profile:
        return {}
    if domain_id:
        data = _profile_domain_data(profile, domain_id)
        if data.get("permissions"):
            return data.get("permissions") or {}
    return profile.permissions or {}


def _get_domain_scope(profile, domain_id=None):
    if not profile:
        return ""
    if domain_id:
        data = _profile_domain_data(profile, domain_id)
        if "scope" in data:
            return (data.get("scope") or "").strip()
    return (profile.scope or "").strip()


def _user_has_domain_access(user, domain_obj) -> bool:
    if not user or not user.is_authenticated or not domain_obj:
        return False
    if _user_is_system_admin(user):
        return True
    profile = _get_user_profile(user)
    if not profile:
        return False
    domain_permissions = profile.domain_permissions or {}
    if domain_permissions:
        data = domain_permissions.get(str(domain_obj.id), {}) or {}
        perms = data.get("permissions") or {}
        return bool(data.get("role") or any(perms.values()))
    return bool(profile.permissions or profile.role in {"admin", "helpdesk", "auditor", "custom"})


def get_available_domains_for_user(user):
    domains = list(LdapSettings.available_settings())
    if not user or not user.is_authenticated:
        return []
    if _user_is_system_admin(user):
        return domains
    return [domain for domain in domains if _user_has_domain_access(user, domain)]


def get_active_ldap_settings(request):
    domains = get_available_domains_for_user(request.user)
    if not domains:
        clear_current_ldap_settings()
        return None
    selected_id = request.session.get("active_ldap_settings_id")
    active = None
    if selected_id:
        active = next((domain for domain in domains if str(domain.id) == str(selected_id)), None)
    if active is None:
        active = domains[0]
        request.session["active_ldap_settings_id"] = active.id
    set_current_ldap_settings(active.id)
    return active


def _domain_context(request):
    if not getattr(request, "user", None) or not request.user.is_authenticated:
        return {"available_domains": [], "active_domain": None}
    available_domains = get_available_domains_for_user(request.user)
    active_domain = get_active_ldap_settings(request) if available_domains else None
    return {
        "available_domains": available_domains,
        "active_domain": active_domain,
    }


def render(request, template_name, context=None, *args, **kwargs):
    context = context or {}
    context.update(_domain_context(request))
    return django_render(request, template_name, context, *args, **kwargs)


def _safe_redirect_target(request, candidate: str | None) -> str | None:
    """Return candidate only when it points back to this host on a safe scheme."""
    if not candidate:
        return None
    if url_has_allowed_host_and_scheme(
        candidate,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return candidate
    return None


@login_required
@require_POST
def select_domain(request):
    domain_id = (request.POST.get("domain_id") or "").strip()
    domains = get_available_domains_for_user(request.user)
    if any(str(domain.id) == str(domain_id) for domain in domains):
        request.session["active_ldap_settings_id"] = int(domain_id)
        set_current_ldap_settings(domain_id)
        messages.success(request, "Active domain changed.")
    else:
        messages.error(request, "You do not have permission to use that domain.")
    safe_next = _safe_redirect_target(request, request.POST.get("next")) or _safe_redirect_target(request, request.META.get("HTTP_REFERER"))
    return redirect(safe_next or "home")

def login_view(request):
    if request.user.is_authenticated:
        return redirect("home")

    if request.method == "POST":
        username = (request.POST.get("username") or "").strip()
        password = request.POST.get("password") or ""

        if not username or not password:
            messages.error(request, "Username and password are required.")
            return render(request, "auth/login.html")

        user = authenticate(request, username=username, password=password)

        if user is None:
            messages.error(request, "Invalid username or password.")
            return render(request, "auth/login.html")

        if not user.is_active:
            messages.error(request, "This account is disabled.")
            return render(request, "auth/login.html")

        login(request, user)
        messages.success(request, f"Welcome back, {user.username}.")

        safe_next = _safe_redirect_target(request, request.GET.get("next"))
        if safe_next:
            return redirect(safe_next)

        return redirect("home")

    return render(request, "auth/login.html")


def logout_view(request):
    logout(request)
    messages.success(request, "You have been signed out.")
    return redirect("login")



@login_required
def index(request):
    default_dashboard = {
        "ad_counts": {"users": 0, "computers": 0, "groups": 0, "ous": 0},
        "user_status": {"enabled": 0, "disabled": 0, "locked": 0, "inactive": 0, "password_never_expires": 0},
        "computer_status": {"enabled": 0, "disabled": 0, "inactive": 0, "domain_controllers": 0},
        "top_ous": [],
        "top_groups": [],
        "computer_os": {"labels": [], "data": []},
        "security_alerts": [],
        "recent_users": [],
        "recent_computers": [],
    }

    active_domain = get_active_ldap_settings(request)

    try:
        from .ad_service import get_dashboard_data
        dashboard = get_dashboard_data()
    except Exception as exc:
        dashboard = default_dashboard
        messages.error(request, f"Failed to load dashboard data: {exc}")

    context = dict(dashboard)
    context["dashboard_json"] = json.dumps(dashboard, default=str)
    return render(request, "index.html", context)


@login_required
@system_admin_required
@require_http_methods(["GET", "POST"])
def ldap_setup(request):
    domain_id = request.POST.get("domain_id") or request.GET.get("domain_id") or request.session.get("active_ldap_settings_id")
    instance = None if request.GET.get("new") == "1" else LdapSettings.get_settings(settings_id=domain_id)
    if request.method == "POST":
        post_domain_id = request.POST.get("domain_id")
        instance = LdapSettings.objects.filter(pk=post_domain_id).first() if post_domain_id else None
        form = LdapSettingsForm(request.POST, instance=instance)
        if "test_only" in request.POST:
            if form.is_valid():
                data = form.cleaned_data
                ok, err = _test_ldap_connection(
                    server_uri=data["server_uri"],
                    use_ssl=data["use_ssl"],
                    bind_dn=data.get("bind_dn") or None,
                    bind_password=data.get("bind_password") or None,
                )
                if ok:
                    messages.success(request, "LDAP connection test: Success.")
                else:
                    messages.error(request, f"LDAP connection test failed: {err}")
            else:
                messages.error(request, "Please fix the errors in the form before testing.")
        else:
            if form.is_valid():
                saved = form.save()
                request.session["active_ldap_settings_id"] = saved.id
                set_current_ldap_settings(saved.id)
                messages.success(request, "LDAP domain settings saved.")
                return redirect(f"{request.path}?domain_id={saved.id}")
            else:
                messages.error(request, "Please fix the errors in the form.")
    else:
        form = LdapSettingsForm(instance=instance)

    return render(request, "ldap_setup.html", {"form": form, "ldap_settings_list": LdapSettings.objects.all().order_by("name", "domain_name", "server_uri"), "current_ldap_setting": instance})



def _test_ldap_connection(server_uri: str, use_ssl: bool, bind_dn: str | None, bind_password: str | None):
    try:
        tls = Tls(validate=ssl.CERT_NONE) if use_ssl else None
        server = Server(server_uri, use_ssl=use_ssl, tls=tls, get_info=ALL)
        conn = Connection(server, user=bind_dn, password=bind_password, auto_bind=True)
        try:
            conn.search(
                search_base="",
                search_filter="(objectClass=*)",
                search_scope=BASE,
                attributes=["defaultNamingContext"],
                size_limit=1,
            )
        except Exception:
            pass
        conn.unbind()
        return True, None
    except Exception as e:
        return False, str(e)


# -----------------------------
# Reports
# -----------------------------

@login_required
@ad_permission_required("reports", "view_reports")
def reports_page(request):
    return render(request, "reports/reports_page.html")



@login_required
@ad_permission_required("reports", "view_reports")
def user_reports_page(request):
    report_type = (request.GET.get("report") or "all").strip().lower()

    report_titles = {
        "all": "All Users",
        "empty_attributes": "Users with Empty Attributes",
        "without_managers": "Users Without Managers",
        "duplicate_attributes": "Users with Duplicate Email",
        "without_email": "Users Without Email",
        "without_hr_id": "Users Without HR ID",
        "enabled": "Enabled Users",
        "recently_created": "Recently Created Users",
        "inactive": "Inactive Users",
        "real_last_logon": "Last Logon Users",
        "recently_logged_on": "Recently Logged On Users",
        "disabled": "Disabled Users",
        "locked": "Locked-out Users",
        "expired": "Account Expired Users",
    }

    report_descriptions = {
        "all": "Showing all available users.",
        "empty_attributes": "Users missing one or more important attributes.",
        "without_managers": "Users who do not have a manager assigned.",
        "duplicate_attributes": "Users with duplicate email values.",
        "without_email": "Users who do not have an email address.",
        "without_hr_id": "Users who do not have an HR ID.",
        "enabled": "Users whose accounts are enabled.",
        "recently_created": "Users created recently based on LDAP creation date.",
        "inactive": "Users with no recent logon activity based on the best available LDAP logon value.",
        "real_last_logon": "Users with available logon values from LDAP attributes.",
        "recently_logged_on": "Users sorted by the most recent available LDAP logon value.",
        "disabled": "Users whose accounts are disabled.",
        "locked": "Users whose accounts appear locked based on LDAP lockout data.",
        "expired": "Users whose accounts are expired.",
    }

    users = []
    try:
        from .ad_service import get_user_report_data
        users = get_user_report_data(report_type=report_type)
        users = sorted(users, key=lambda x: (x.get("display_name") or x.get("username") or "").lower())
    except Exception as e:
        messages.error(request, f"Failed to load user report data: {e}")

    context = {
        "users": users,
        "selected_report": report_type,
        "report_title": report_titles.get(report_type, "User Report"),
        "report_description": report_descriptions.get(report_type, "Showing user report data."),
    }
    return render(request, "reports/user/user_reports_page.html", context)



@login_required
@ad_permission_required("reports", "view_reports")
def group_reports_page(request):
    report_type = (request.GET.get("report") or "all").strip().lower()

    report_titles = {
        "all": "All Groups",
        "with_members": "Groups With Members",
        "detailed_members": "Detailed Group Members",
        "without_members": "Groups Without Members",
        "nested_groups": "Nested Groups",
        "recently_created": "Recently Created Groups",
        "recently_modified": "Recently Modified Groups",
        "without_description": "Groups Without Description",
        "security": "Security Groups",
        "distribution": "Distribution Groups",
        "large_groups": "Large Groups",
    }

    report_descriptions = {
        "all": "Showing all available groups.",
        "with_members": "Groups that currently have one or more members.",
        "detailed_members": "Groups with member preview information.",
        "without_members": "Groups that do not currently contain any members.",
        "nested_groups": "Groups that contain one or more nested groups.",
        "recently_created": "Groups created recently based on LDAP creation date.",
        "recently_modified": "Groups modified recently based on LDAP change date.",
        "without_description": "Groups without a description.",
        "security": "Groups marked as security groups.",
        "distribution": "Groups marked as distribution groups.",
        "large_groups": "Groups with a large number of members.",
    }

    groups = []
    try:
        from .ad_service import get_group_report_data
        groups = get_group_report_data(report_type=report_type)
        groups = sorted(groups, key=lambda x: (x.get("name") or "").lower())
    except Exception as e:
        messages.error(request, f"Failed to load group report data: {e}")

    context = {
        "groups": groups,
        "selected_report": report_type,
        "report_title": report_titles.get(report_type, "Group Report"),
        "report_description": report_descriptions.get(report_type, "Showing group report data."),
    }
    return render(request, "reports/group/group_reports_page.html", context)



@login_required
@ad_permission_required("reports", "view_reports")
def computer_reports_page(request):
    report_type = (request.GET.get("report") or "all").strip().lower()

    report_titles = {
        "all": "All Computers",
        "os_based": "OS Based Report",
        "workstations": "Workstation Computers",
        "inactive": "Inactive Computers",
        "active": "Active Computers",
        "disabled": "Disabled Computers",
        "bitlocker_keys": "BitLocker Recovery Keys",
        "bitlocker_enabled": "BitLocker Enabled Computers",
    }

    report_descriptions = {
        "all": "Showing all available computer accounts.",
        "os_based": "Computers grouped or filtered by available operating system data.",
        "workstations": "Computers that appear to be workstation devices based on operating system data.",
        "inactive": "Computers with no recent logon activity based on the best available LDAP logon value.",
        "active": "Computer accounts that are currently enabled.",
        "disabled": "Computer accounts that are currently disabled.",
        "bitlocker_keys": "Computers with BitLocker recovery keys available in Active Directory.",
        "bitlocker_enabled": "Computers with BitLocker recovery information available in Active Directory.",
    }

    computers = []
    try:
        computers = get_computer_report_data(report_type=report_type)
        computers = sorted(computers, key=lambda x: (x.get("name") or x.get("sam") or "").lower())
    except Exception as e:
        messages.error(request, f"Failed to load computer report data: {e}")

    context = {
        "computers": computers,
        "selected_report": report_type,
        "report_title": report_titles.get(report_type, "Computer Report"),
        "report_description": report_descriptions.get(report_type, "Showing computer report data."),
    }
    return render(request, "reports/computer/computer_reports_page.html", context)



@login_required
@ad_permission_required("reports", "view_reports")
def ou_reports_page(request):
    report_type = (request.GET.get("report") or "all").strip().lower()

    report_titles = {
        "all": "All OUs",
        "empty": "Empty OUs",
        "protected": "Protected OUs",
        "unprotected": "Unprotected OUs",
        "recently_created": "Recently Created OUs",
        "recently_modified": "Recently Modified OUs",
        "counts": "OU Object Counts",
        "gpo_linked": "GPO-linked OUs",
    }

    report_descriptions = {
        "all": "Showing all available organizational units.",
        "empty": "OUs with no immediate users, computers, groups, or child OUs.",
        "protected": "OUs marked as protected from accidental deletion when WinRM/AD PowerShell data is available.",
        "unprotected": "OUs not marked as protected from accidental deletion when WinRM/AD PowerShell data is available.",
        "recently_created": "OUs created recently based on LDAP creation date.",
        "recently_modified": "OUs modified recently based on LDAP changed date.",
        "counts": "Immediate object counts for each OU.",
        "gpo_linked": "OUs with one or more linked Group Policy Objects.",
    }

    ous = []
    try:
        ous = get_ou_report_data(report_type=report_type)
        ous = sorted(ous, key=lambda x: (x.get("name") or x.get("ou") or "").lower())
    except Exception as e:
        messages.error(request, f"Failed to load OU report data: {e}")

    context = {
        "ous": ous,
        "report": report_type,
        "selected_report": report_type,
        "table_title": report_titles.get(report_type, "OU Report"),
        "report_title": report_titles.get(report_type, "OU Report"),
        "report_description": report_descriptions.get(report_type, "Showing OU report data."),
    }
    return render(request, "reports/ou/ou_reports_page.html", context)


# -----------------------------
# User management
# -----------------------------

@login_required
@ad_permission_required("users", "any")
def user_management(request):
    return render(request, "management/user/user_management.html")



@login_required
@ad_permission_required("users", "unlock")
def unlock_single_user(request):
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        try:
            unlock_user(username)
            messages.success(request, f"User unlocked: {username}")
        except ADServiceError as e:
            messages.error(request, f"Unlock failed: {e}")
    return render(request, "management/user/unlock/single_user.html")



@login_required
@ad_permission_required("users", "unlock")
def unlock_bulk_users(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                username = str(row.get("username", "") or "").strip()
                if username:
                    unlock_user(username)
                    count += 1
            messages.success(request, f"Bulk unlock completed. {count} users unlocked.")
        except ADServiceError as e:
            messages.error(request, f"Bulk unlock failed: {e}")
    return render(request, "management/user/unlock/bulk_users.html")



@login_required
@ad_permission_required("users", "reset")
def reset_single_user(request):
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        new_password = request.POST.get("new_password", "")
        confirm_password = request.POST.get("confirm_password", "")
        if new_password != confirm_password:
            messages.error(request, "Passwords do not match.")
        else:
            try:
                reset_user_password(username, new_password)
                messages.success(request, f"Password reset for {username}.")
            except ADServiceError as e:
                messages.error(request, f"Password reset failed: {e}")
    return render(request, "management/user/reset/single_user.html")



@login_required
@ad_permission_required("users", "reset")
def reset_bulk_users(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                username = str(row.get("username", "") or "").strip()
                new_password = str(row.get("new_password", "") or "").strip()
                if username and new_password:
                    reset_user_password(username, new_password)
                    count += 1
            messages.success(request, f"Bulk password reset completed. {count} users updated.")
        except ADServiceError as e:
            messages.error(request, f"Bulk reset failed: {e}")
    return render(request, "management/user/reset/bulk_users.html")



@login_required
@ad_permission_required("users", "lock")
def lock_single_user(request):
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        try:
            lock_user(username)
            messages.success(request, f"User locked: {username}")
        except ADServiceError as e:
            messages.error(request, f"Lock failed: {e}")
    return render(request, "management/user/lock/single_user.html")



@login_required
@ad_permission_required("users", "lock")
def lock_bulk_users(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                username = str(row.get("username", "") or "").strip()
                if username:
                    lock_user(username)
                    count += 1
            messages.success(request, f"Bulk lock completed. {count} users locked.")
        except ADServiceError as e:
            messages.error(request, f"Bulk lock failed: {e}")
    return render(request, "management/user/lock/bulk_users.html")



@login_required
@ad_permission_required("users", "modify")
def update_single_user(request):
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        updates = {
            "givenName": request.POST.get("first_name", "").strip(),
            "sn": request.POST.get("last_name", "").strip(),
            "displayName": request.POST.get("display_name", "").strip(),
            "mail": request.POST.get("email", "").strip(),
            "telephoneNumber": request.POST.get("phone", "").strip(),
            "department": request.POST.get("department", "").strip(),
            "description": request.POST.get("description", "").strip(),
            "employeeID": request.POST.get("hr_id", "").strip(),
        }

        try:
            update_user(username, updates)
            messages.success(request, f"User updated: {username}")
        except ADServiceError as e:
            messages.error(request, f"Update failed: {e}")
        except Exception as e:
            messages.error(request, f"Unexpected error: {e}")

    return render(request, "management/user/update/single_user.html")



@login_required
@ad_permission_required("users", "modify")
def update_bulk_users(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)

            updated_count = 0
            errors = []

            for index, row in enumerate(rows, start=2):
                try:
                    username = str(row.get("username", "") or "").strip()
                    if not username:
                        errors.append(f"Row {index}: username is required.")
                        continue

                    ou_name = str(row.get("ou_name", "") or "").strip()
                    target_ou_dn = resolve_ou_name_to_dn(ou_name) if ou_name else None

                    updates = {
                        "givenName": str(row.get("first_name", "") or "").strip(),
                        "sn": str(row.get("last_name", "") or "").strip(),
                        "displayName": str(row.get("display_name", "") or "").strip(),
                        "mail": str(row.get("email", "") or "").strip(),
                        "telephoneNumber": str(row.get("phone", "") or "").strip(),
                        "department": str(row.get("department", "") or "").strip(),
                        "description": str(row.get("description", "") or "").strip(),
                        "employeeID": str(row.get("hr_id", "") or "").strip(),
                    }

                    update_user(username, updates)

                    if target_ou_dn:
                        move_user(username, target_ou_dn)

                    updated_count += 1

                except ADServiceError as e:
                    errors.append(f"Row {index}: {e}")
                except Exception as e:
                    errors.append(f"Row {index}: Unexpected error: {e}")

            if updated_count:
                messages.success(request, f"{updated_count} users updated successfully.")

            if errors:
                for err in errors[:20]:
                    messages.error(request, err)
                if len(errors) > 20:
                    messages.error(request, f"And {len(errors) - 20} more errors.")

        except ADServiceError as e:
            messages.error(request, f"Bulk update failed: {e}")
        except Exception as e:
            messages.error(request, f"Unexpected error: {e}")

    return render(request, "management/user/update/bulk_users.html")



@login_required
@ad_permission_required("users", "create", scope_keys=["target_ou_dn"])
def create_single_user(request):
    try:
        ous = list_ous()
    except Exception:
        ous = []

    try:
        groups = list_groups()
    except Exception:
        groups = []

    if request.method == "POST":
        first_name = request.POST.get("first_name", "").strip()
        last_name = request.POST.get("last_name", "").strip()
        username = request.POST.get("username", "").strip()
        email = request.POST.get("email", "").strip()
        phone = request.POST.get("phone", "").strip()
        department = request.POST.get("department", "").strip()
        description = request.POST.get("description", "").strip()
        hr_id = request.POST.get("hr_id", "").strip()
        password = request.POST.get("password", "")
        confirm_pass = request.POST.get("confirm_password", "")
        must_change_password = bool(request.POST.get("must_change_password"))
        user_cannot_change_password = bool(request.POST.get("user_cannot_change_password"))
        password_never_expires = bool(request.POST.get("password_never_expires"))
        account_disabled = bool(request.POST.get("account_disabled"))

        target_ou_dn = request.POST.get("target_ou_dn", "").strip()
        group_dns = request.POST.getlist("group_dns")

        if not username or not password:
            messages.error(request, "Username and password are required.")
            return render(request, "management/user/create/single_user.html", {"ous": ous, "groups": groups})

        if password != confirm_pass:
            messages.error(request, "Passwords do not match.")
            return render(request, "management/user/create/single_user.html", {"ous": ous, "groups": groups})

        if not target_ou_dn:
            messages.error(request, "Please select the target OU.")
            return render(request, "management/user/create/single_user.html", {"ous": ous, "groups": groups})

        try:
            create_user(
                username=username,
                first_name=first_name,
                last_name=last_name,
                email=email,
                password=password,
                phone=phone,
                department=department,
                description=description,
                hr_id=hr_id,
                target_ou_dn=target_ou_dn or None,
                group_dns=group_dns or None,
                must_change_password=must_change_password,
                user_cannot_change_password=user_cannot_change_password,
                password_never_expires=password_never_expires,
                account_disabled=account_disabled,
            )
            messages.success(request, f"User created in AD: {username}")
        except ADServiceError as e:
            messages.error(request, f"Create failed: {e}")
        except Exception as e:
            messages.error(request, f"Unexpected error: {e}")

    return render(request, "management/user/create/single_user.html", {"ous": ous, "groups": groups})



@login_required
@ad_permission_required("users", "create")
def create_bulk_users(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)

            _validate_required_headers(
                rows,
                ["username", "first_name", "last_name", "email", "password", "ou_name"],
            )

            created_count = 0
            errors = []

            for index, row in enumerate(rows, start=2):
                try:
                    username = str(row.get("username", "") or "").strip()
                    first_name = str(row.get("first_name", "") or "").strip()
                    last_name = str(row.get("last_name", "") or "").strip()
                    email = str(row.get("email", "") or "").strip()
                    password = str(row.get("password", "") or "").strip()
                    phone = str(row.get("phone", "") or "").strip()
                    department = str(row.get("department", "") or "").strip()
                    description = str(row.get("description", "") or "").strip()
                    hr_id = str(row.get("hr_id", "") or "").strip()
                    ou_name = str(row.get("ou_name", "") or "").strip()

                    if not username:
                        errors.append(f"Row {index}: username is required.")
                        continue

                    if not password:
                        errors.append(f"Row {index}: password is required.")
                        continue

                    target_ou_dn = resolve_ou_name_to_dn(ou_name) if ou_name else None

                    create_user(
                        username=username,
                        first_name=first_name,
                        last_name=last_name,
                        email=email,
                        password=password,
                        phone=phone,
                        department=department,
                        description=description,
                        hr_id=hr_id,
                        target_ou_dn=target_ou_dn,
                    )
                    created_count += 1

                except ADServiceError as e:
                    errors.append(f"Row {index}: {e}")
                except Exception as e:
                    errors.append(f"Row {index}: Unexpected error: {e}")

            if created_count:
                messages.success(request, f"{created_count} users created successfully.")

            if errors:
                for err in errors[:20]:
                    messages.error(request, err)
                if len(errors) > 20:
                    messages.error(request, f"And {len(errors) - 20} more errors.")

        except ADServiceError as e:
            messages.error(request, f"Bulk create failed: {e}")
        except Exception as e:
            messages.error(request, f"Unexpected error: {e}")

    return render(request, "management/user/create/bulk_users.html")



@login_required
@ad_permission_required("users", "modify", scope_keys=["target_ou_dn"])
def move_single_user(request):
    try:
        ous = list_ous()
    except Exception:
        ous = []

    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        target_ou_dn = request.POST.get("target_ou_dn", "").strip()

        try:
            move_user(username, target_ou_dn)
            messages.success(request, f"User moved successfully: {username}")
        except ADServiceError as e:
            messages.error(request, f"Move failed: {e}")

    return render(request, "management/user/move/single_user.html", {"ous": ous})



@login_required
@ad_permission_required("users", "modify")
def move_bulk_users(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            moved_count = 0
            errors = []

            for index, row in enumerate(rows, start=2):
                try:
                    username = str(row.get("username", "") or "").strip()
                    ou_name = str(row.get("ou_name", "") or "").strip()
                    if not username or not ou_name:
                        errors.append(f"Row {index}: username and ou_name are required.")
                        continue

                    target_ou_dn = resolve_ou_name_to_dn(ou_name)
                    move_user(username, target_ou_dn)
                    moved_count += 1
                except ADServiceError as e:
                    errors.append(f"Row {index}: {e}")

            if moved_count:
                messages.success(request, f"Bulk move completed. {moved_count} users moved.")
            if errors:
                for err in errors[:20]:
                    messages.error(request, err)

        except ADServiceError as e:
            messages.error(request, f"Bulk move failed: {e}")
    return render(request, "management/user/move/bulk_users.html")


# -----------------------------
# Computer management
# -----------------------------

@login_required
@ad_permission_required("computers", "any")
def computer_management(request):
    return render(request, "management/computer/computer_management_page.html")



@login_required
@ad_permission_required("computers", "create", scope_keys=["ou"])
def create_single_computer(request):
    ous = []
    ou_error = None

    try:
        ous = list_ous()
    except Exception as exc:
        ou_error = str(exc)

    if request.method == "POST":
        computer_name = request.POST.get("computer_name", "").strip()
        ou = request.POST.get("ou", "").strip()
        description = request.POST.get("description", "").strip()
        try:
            if not computer_name:
                messages.error(request, "Computer name is required.")
            else:
                try:
                    create_computer(computer_name, ou or None, description)
                except TypeError:
                    create_computer(computer_name, ou or None)
                messages.success(request, f"Computer created: {computer_name}")
        except ADServiceError as e:
            messages.error(request, f"Create failed: {e}")

    return render(request, "management/computer/create/single_computer.html", {"ous": ous, "ou_error": ou_error})



@login_required
@ad_permission_required("computers", "create")
def create_bulk_computers(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                name = str(row.get("computer_name", "") or "").strip() or str(row.get("name", "") or "").strip()
                ou = str(row.get("ou", "") or "").strip()
                description = str(row.get("description", "") or "").strip()
                if name:
                    try:
                        create_computer(name, ou or None, description)
                    except TypeError:
                        create_computer(name, ou or None)
                    count += 1
            messages.success(request, f"Bulk computer creation completed. {count} computers created.")
        except ADServiceError as e:
            messages.error(request, f"Bulk create failed: {e}")
    return render(request, "management/computer/create/bulk_computers.html")



@login_required
@ad_permission_required("computers", "lock")
def lock_single_computer(request):
    if request.method == "POST":
        computer_name = request.POST.get("computer_name", "").strip()
        try:
            lock_computer(computer_name)
            messages.success(request, f"Computer locked: {computer_name}")
        except ADServiceError as e:
            messages.error(request, f"Lock failed: {e}")
    return render(request, "management/computer/lock/single_computer.html")



@login_required
@ad_permission_required("computers", "lock")
def lock_bulk_computers(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                name = str(row.get("computer_name", "") or "").strip() or str(row.get("name", "") or "").strip()
                if name:
                    lock_computer(name)
                    count += 1
            messages.success(request, f"Bulk lock completed. {count} computers locked.")
        except ADServiceError as e:
            messages.error(request, f"Bulk lock failed: {e}")
    return render(request, "management/computer/lock/bulk_computers.html")



@login_required
@ad_permission_required("computers", "unlock")
def unlock_single_computer(request):
    if request.method == "POST":
        computer_name = request.POST.get("computer_name", "").strip()
        try:
            unlock_computer(computer_name)
            messages.success(request, f"Computer unlocked: {computer_name}")
        except ADServiceError as e:
            messages.error(request, f"Unlock failed: {e}")
    return render(request, "management/computer/unlock/single_computer.html")



@login_required
@ad_permission_required("computers", "unlock")
def unlock_bulk_computers(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                name = str(row.get("computer_name", "") or "").strip() or str(row.get("name", "") or "").strip()
                if name:
                    unlock_computer(name)
                    count += 1
            messages.success(request, f"Bulk unlock completed. {count} computers unlocked.")
        except ADServiceError as e:
            messages.error(request, f"Bulk unlock failed: {e}")
    return render(request, "management/computer/unlock/bulk_computers.html")



@login_required
@ad_permission_required("computers", "modify", scope_keys=["target_ou"])
def move_single_computer(request):
    ous = []
    ou_error = None

    try:
        ous = list_ous()
    except Exception as exc:
        ou_error = str(exc)

    if request.method == "POST":
        computer_name = request.POST.get("computer_name", "").strip()
        target_ou = request.POST.get("target_ou", "").strip()

        if not computer_name:
            messages.error(request, "Computer name is required.")
            return render(request, "management/computer/move/single_computer.html", {"ous": ous, "ou_error": ou_error})

        if not target_ou:
            messages.error(request, "Target OU is required.")
            return render(request, "management/computer/move/single_computer.html", {"ous": ous, "ou_error": ou_error})

        try:
            move_computer(computer_name, target_ou)
            messages.success(request, f"Computer '{computer_name}' moved successfully.")
            return redirect("move_single_computer")
        except ADServiceError as exc:
            messages.error(request, f"Failed to move computer '{computer_name}': {exc}")
        except Exception as exc:
            messages.error(request, f"Failed to move computer '{computer_name}': {exc}")

    return render(request, "management/computer/move/single_computer.html", {"ous": ous, "ou_error": ou_error})



@login_required
@ad_permission_required("computers", "modify")
def move_bulk_computers(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                name = str(row.get("computer_name", "") or "").strip() or str(row.get("name", "") or "").strip()
                target_ou = str(row.get("target_ou", "") or row.get("target_ou_dn", "") or "").strip()
                if name and target_ou:
                    move_computer(name, target_ou)
                    count += 1
            messages.success(request, f"Bulk move completed. {count} computers moved.")
        except ADServiceError as e:
            messages.error(request, f"Bulk move failed: {e}")
    return render(request, "management/computer/move/bulk_computers.html")



@login_required
@ad_permission_required("computers", "modify")
def update_single_computer(request):
    if request.method == "POST":
        computer_name = request.POST.get("computer_name", "").strip()
        updates = {
            "description": request.POST.get("description", "").strip(),
        }
        try:
            update_computer(computer_name, updates)
            messages.success(request, f"Computer updated: {computer_name}")
        except ADServiceError as e:
            messages.error(request, f"Update failed: {e}")
    return render(request, "management/computer/update/single_computer.html")



@login_required
@ad_permission_required("computers", "modify")
def update_bulk_computers(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                name = str(row.get("computer_name", "") or "").strip() or str(row.get("name", "") or "").strip()
                if not name:
                    continue
                updates = {"description": str(row.get("description", "") or "").strip()}
                update_computer(name, updates)
                count += 1
            messages.success(request, f"Bulk update completed. {count} computers updated.")
        except ADServiceError as e:
            messages.error(request, f"Bulk update failed: {e}")
    return render(request, "management/computer/update/bulk_computers.html")


# -----------------------------
# OU management
# -----------------------------

@login_required
@ad_permission_required("ous", "any")
def ou_management(request):
    return render(request, "management/ou/ou_management.html")



@login_required
@ad_permission_required("ous", "create", scope_keys=["parent_ou"])
def create_single_ou(request):
    if request.method == "POST":
        ou_name = request.POST.get("ou_name", "").strip()
        parent_ou = request.POST.get("parent_ou", "").strip()
        description = request.POST.get("description", "").strip()
        protect = bool(request.POST.get("protect"))

        if not ou_name:
            messages.error(request, "OU Name is required.")
            return redirect("create_single_ou")

        try:
            create_ou(ou_name, parent_ou or None, description, protect)
            messages.success(request, f"OU '{ou_name}' created successfully.")
        except ADServiceError as e:
            messages.error(request, f"Failed to create OU '{ou_name}': {e}")
        except Exception as e:
            messages.error(request, f"Unexpected error while creating OU '{ou_name}': {e}")

        return redirect("create_single_ou")

    return render(request, "management/ou/create_single_ou.html")



@login_required
@ad_permission_required("ous", "modify", scope_keys=["ou_dn"])
def update_single_ou(request):
    ous = []
    selected_ou = None
    ou_details = None
    ou_error = None

    try:
        ous = list_ous()
    except Exception as exc:
        ou_error = str(exc)

    if request.method == "POST":
        selected_ou = request.POST.get("ou_dn", "").strip()
        new_name = request.POST.get("new_name", "").strip()
        description = request.POST.get("description", "").strip()
        protect = bool(request.POST.get("protect"))

        if not selected_ou:
            messages.error(request, "Please select an OU.")
            return render(request, "management/ou/update_single_ou.html", {
                "ous": ous,
                "ou_error": ou_error,
                "selected_ou": selected_ou,
                "ou_details": ou_details,
            })

        try:
            update_ou(selected_ou, new_name or None, description, protect)
            messages.success(request, "OU updated successfully.")
            return redirect("update_single_ou")
        except Exception as exc:
            messages.error(request, f"Failed to update OU: {exc}")

    else:
        selected_ou = request.GET.get("ou_dn", "").strip()
        if selected_ou:
            try:
                from .ad_service import get_ou_details
                ou_details = get_ou_details(selected_ou)
            except Exception as exc:
                messages.error(request, f"Failed to load OU details: {exc}")

    return render(request, "management/ou/update_single_ou.html", {
        "ous": ous,
        "ou_error": ou_error,
        "selected_ou": selected_ou,
        "ou_details": ou_details,
    })



@login_required
@ad_permission_required("ous", "modify", scope_keys=["target_parent_dn", "ou_dn"])
def move_ou(request):
    ous = []
    ou_error = None

    try:
        ous = list_ous()
    except Exception as exc:
        ou_error = str(exc)

    if request.method == "POST":
        ou_dn = request.POST.get("ou_dn", "").strip()
        target_parent_dn = request.POST.get("target_parent_dn", "").strip()

        if not ou_dn or not target_parent_dn:
            messages.error(request, "Please select both source and target OU.")
        else:
            try:
                ad_move_ou(ou_dn, target_parent_dn)
                messages.success(request, "OU moved successfully.")
                return redirect("move_ou")
            except Exception as exc:
                messages.error(request, f"Failed to move OU: {exc}")

    return render(request, "management/ou/move_ou.html", {"ous": ous, "ou_error": ou_error})



@login_required
@ad_permission_required("ous", "delete", scope_keys=["ou_dn"])
def delete_ou(request):
    ou_error = None
    try:
        ous = list_ous()
    except Exception as exc:
        ous = []
        ou_error = str(exc)

    ctx = {"ous": ous, "ou_error": ou_error}

    if request.method == "POST":
        ou_dn = request.POST.get("ou_dn", "").strip()
        include_children = bool(request.POST.get("include_children"))
        dry_run = bool(request.POST.get("dry_run"))
        confirm = bool(request.POST.get("confirm"))

        if not ou_dn:
            messages.error(request, "Please provide the OU Distinguished Name (DN).")
            return render(request, "management/ou/delete_ou.html", ctx)

        if not confirm:
            messages.error(request, "You must confirm that you understand this action is irreversible.")
            return render(request, "management/ou/delete_ou.html", ctx)

        if dry_run:
            messages.info(request, f"Dry run only. Would delete OU: {ou_dn} (include children: {include_children}).")
            return render(request, "management/ou/delete_ou.html", ctx)

        try:
            ad_delete_ou(ou_dn)
            messages.success(request, f"OU deleted successfully: {ou_dn}")
            return redirect("ou_management")
        except ADServiceError as e:
            messages.error(request, f"Delete failed: {e}")

    return render(request, "management/ou/delete_ou.html", ctx)



@login_required
@ad_permission_required("ous", "create")
def create_bulk_ous(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                ou_name = str(row.get("ou_name", "") or "").strip()
                parent_ou = str(row.get("parent_ou", "") or "").strip()
                description = str(row.get("description", "") or "").strip()
                protect = str(row.get("protect", "") or "").strip().lower() in ("1", "true", "yes", "y")
                if ou_name:
                    create_ou(ou_name, parent_ou or None, description, protect)
                    count += 1
            messages.success(request, f"Bulk OU create completed. {count} OUs created.")
        except ADServiceError as e:
            messages.error(request, f"Bulk create failed: {e}")
    return render(request, "management/ou/create_bulk_ou.html")



@login_required
@ad_permission_required("ous", "create")
def download_bulk_ou_sample(request):
    file_path = os.path.join(settings.BASE_DIR, "static", "samples", "bulk_ou_sample.xlsx")

    if not os.path.exists(file_path):
        raise Http404("Sample file not found.")

    return FileResponse(open(file_path, "rb"), as_attachment=True, filename="bulk_ou_sample.xlsx")



@login_required
@ad_permission_required("ous", "modify")
def update_bulk_ous(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                ou_dn = str(row.get("ou_dn", "") or "").strip()
                new_name = str(row.get("new_name", "") or "").strip()
                description = str(row.get("description", "") or "").strip()
                protect = str(row.get("protect", "") or "").strip().lower() in ("1", "true", "yes", "y")
                if ou_dn:
                    update_ou(ou_dn, new_name, description, protect)
                    count += 1
            messages.success(request, f"Bulk OU update completed. {count} OUs updated.")
        except ADServiceError as e:
            messages.error(request, f"Bulk update failed: {e}")
    return render(request, "management/ou/update_bulk_ou.html")



@login_required
@ad_permission_required("ous", "delete")
def delete_bulk_ous(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                ou_dn = str(row.get("ou_dn", "") or "").strip()
                if ou_dn:
                    ad_delete_ou(ou_dn)
                    count += 1
            messages.success(request, f"Bulk OU delete completed. {count} OUs deleted.")
        except ADServiceError as e:
            messages.error(request, f"Bulk delete failed: {e}")
    return render(request, "management/ou/delete_bulk_ou.html")



@login_required
@ad_permission_required("ous", "modify")
def move_bulk_ous(request):
    if request.method == "POST":
        file = _get_uploaded_file(request, "file", "excel_file")
        try:
            rows = _read_excel_rows(file)
            count = 0
            for row in rows:
                ou_dn = str(row.get("ou_dn", "") or "").strip()
                target_parent = str(row.get("target_parent_dn", "") or "").strip()
                if ou_dn and target_parent:
                    ad_move_ou(ou_dn, target_parent)
                    count += 1
            messages.success(request, f"Bulk OU move completed. {count} OUs moved.")
        except ADServiceError as e:
            messages.error(request, f"Bulk move failed: {e}")
    return render(request, "management/ou/move_bulk_ou.html")


# -----------------------------
# Group management
# -----------------------------

@login_required
@ad_permission_required("groups", "any")
def group_management(request):
    return render(request, "management/group/group_management_page.html")



@login_required
@ad_permission_required("groups", "create", scope_keys=["ou_dn"])
def create_single_group(request):
    ous = list_group_ous()

    if request.method == "POST":
        group_name = request.POST.get("group_name", "").strip()
        description = request.POST.get("description", "").strip()

        user_dns = [u.strip() for u in request.POST.getlist("user_dns") if u.strip()]
        computer_dns = [c.strip() for c in request.POST.getlist("computer_dns") if c.strip()]
        nested_group_dns = [g.strip() for g in request.POST.getlist("nested_group_dns") if g.strip()]

        group_scope = request.POST.get("group_scope", "Global").strip() or "Global"
        group_category = request.POST.get("group_category", "Security").strip() or "Security"
        owner_dn = request.POST.get("owner_dn", "").strip()
        ou_dn = request.POST.get("ou_dn", "").strip()
        copy_from_group_dn = request.POST.get("copy_from_group_dn", "").strip()
        protect_from_deletion = request.POST.get("protect_from_deletion") == "1"

        if not group_name:
            messages.error(request, "Group name is required.")
        else:
            try:
                create_group(
                    group_name=group_name,
                    description=description,
                    user_dns=user_dns,
                    computer_dns=computer_dns,
                    nested_group_dns=nested_group_dns,
                    group_scope=group_scope,
                    group_category=group_category,
                    owner_dn=owner_dn,
                    ou_dn=ou_dn,
                    protect_from_deletion=protect_from_deletion,
                    copy_from_group_dn=copy_from_group_dn,
                    skip_invalid_members=False,
                )
                messages.success(request, f"Group created successfully: {group_name}")
            except ADServiceError as e:
                messages.error(request, f"Group creation failed: {e}")
            except Exception as e:
                messages.error(request, f"Unexpected error while creating group: {e}")

    context = {
        "ous": ous,
        "scope_choices": [
            ("Global", "Global"),
            ("DomainLocal", "Domain Local"),
            ("Universal", "Universal"),
        ],
        "category_choices": [
            ("Security", "Security"),
            ("Distribution", "Distribution"),
        ],
    }
    return render(request, "management/group/create/single_group.html", context)



@login_required
@ad_permission_required("groups", "create")
def create_bulk_groups(request):
    results = []

    if request.method == "POST":
        upload = _get_uploaded_file(request, "file", "excel_file")

        if not upload:
            messages.error(request, "Please upload an Excel file.")
        else:
            filename = (upload.name or "").lower()
            if not filename.endswith(".xlsx"):
                messages.error(request, "Please upload a .xlsx Excel file.")
            else:
                try:
                    results = bulk_create_groups_from_excel(upload)

                    success_count = sum(1 for row in results if row.get("success"))
                    fail_count = len(results) - success_count

                    if success_count:
                        messages.success(request, f"{success_count} group(s) created successfully.")
                    if fail_count:
                        messages.warning(request, f"{fail_count} row(s) failed. Review the result table below.")
                    if not results:
                        messages.info(request, "The uploaded file did not contain any data rows.")
                except ADServiceError as e:
                    messages.error(request, f"Bulk group creation failed: {e}")
                except Exception as e:
                    messages.error(request, f"Unexpected error during bulk group creation: {e}")

    return render(request, "management/group/create/bulk_group.html", {"results": results})



@login_required
def search_users_view(request):
    query = request.GET.get("q", "").strip()
    limit = int(request.GET.get("limit", 20))
    results = search_users(query, limit=limit) if query else []
    return JsonResponse({"results": results})



@login_required
def search_computers_view(request):
    query = request.GET.get("q", "").strip()
    limit = int(request.GET.get("limit", 20))
    results = search_computers(query, limit=limit) if query else []
    return JsonResponse({"results": results})



@login_required
def search_groups_view(request):
    query = request.GET.get("q", "").strip()
    limit = int(request.GET.get("limit", 20))
    results = search_groups(query, limit=limit) if query else []
    return JsonResponse({"results": results})



@login_required
@ad_permission_required("groups", "modify")
def update_single_group(request):
    if request.method == "POST":
        group_name = request.POST.get("group_name", "").strip()
        description = request.POST.get("description", "").strip()
        group_scope = request.POST.get("group_scope", "").strip()
        group_category = request.POST.get("group_category", "").strip()

        add_member_dns = [v.strip() for v in request.POST.getlist("add_member_dns") if v.strip()]
        remove_member_dns = [v.strip() for v in request.POST.getlist("remove_member_dns") if v.strip()]

        if not group_name:
            messages.error(request, "Group name is required.")
        else:
            try:
                update_group(
                    group_name=group_name,
                    description=description,
                    group_scope=group_scope,
                    group_category=group_category,
                    add_member_dns=add_member_dns,
                    remove_member_dns=remove_member_dns,
                )
                messages.success(request, f"Group updated successfully: {group_name}")
            except ADServiceError as e:
                messages.error(request, f"Group update failed: {e}")
            except Exception as e:
                messages.error(request, f"Unexpected error: {e}")

    return render(request, "management/group/update/single_group.html")



@login_required
@ad_permission_required("groups", "modify")
def update_bulk_groups(request):
    results = []

    if request.method == "POST":
        upload = _get_uploaded_file(request, "file", "excel_file")

        if not upload:
            messages.error(request, "Please upload an Excel file.")
        else:
            try:
                results = bulk_update_groups_from_excel(upload)
                success_count = sum(1 for r in results if r["success"])
                fail_count = len(results) - success_count

                if success_count:
                    messages.success(request, f"{success_count} group rows updated successfully.")
                if fail_count:
                    messages.warning(request, f"{fail_count} rows failed.")
            except ADServiceError as e:
                messages.error(request, f"Bulk group update failed: {e}")
            except Exception as e:
                messages.error(request, f"Unexpected error: {e}")

    return render(request, "management/group/update/bulk_group.html", {"results": results})



@login_required
def search_directory_objects_ajax(request):
    query = request.GET.get("q", "").strip()
    results = search_directory_objects(query) if query else []
    return JsonResponse({"results": results})



@login_required
def get_group_members_ajax(request):
    group_name = request.GET.get("group_name", "").strip()
    if not group_name:
        return JsonResponse({"results": []})

    try:
        results = get_group_members_for_ui(group_name)
        return JsonResponse({"results": results})
    except ADServiceError as e:
        return JsonResponse({"results": [], "error": str(e)}, status=400)



@login_required
@ad_permission_required("groups", "delete")
def delete_single_group(request):
    if request.method == "POST":
        group_name = request.POST.get("group_name", "").strip()
        confirm = request.POST.get("confirm_delete")
        if not confirm:
            messages.error(request, "Please confirm deletion.")
        else:
            try:
                delete_group(group_name)
                messages.success(request, f"Group deleted: {group_name}")
            except ADServiceError as e:
                messages.error(request, f"Delete failed: {e}")
    return render(request, "management/group/delete/single_group.html")



@login_required
@ad_permission_required("groups", "modify", scope_keys=["target_ou_dn"])
def move_single_group(request):
    try:
        ous = list_group_ous()
    except Exception:
        ous = []

    if request.method == "POST":
        group_name = request.POST.get("group_name", "").strip()
        target_ou_dn = request.POST.get("target_ou_dn", "").strip()

        try:
            move_group(group_name, target_ou_dn)
            messages.success(request, f"Group moved successfully: {group_name}")
        except ADServiceError as e:
            messages.error(request, f"Move failed: {e}")

    return render(request, "management/group/move/single_group.html", {"ous": ous})



@login_required
@ad_permission_required("groups", "modify")
def move_bulk_groups(request):
    results = []

    if request.method == "POST":
        upload = _get_uploaded_file(request, "file", "excel_file")

        if not upload:
            messages.error(request, "Please upload an Excel file.")
        else:
            try:
                results = bulk_move_groups_from_excel(upload)

                success_count = sum(1 for r in results if r["success"])
                fail_count = len(results) - success_count

                if success_count:
                    messages.success(request, f"{success_count} groups moved successfully.")
                if fail_count:
                    messages.warning(request, f"{fail_count} rows failed.")
            except Exception as e:
                messages.error(request, f"Bulk move failed: {e}")

    return render(request, "management/group/move/bulk_group.html", {"results": results, "mode": "move"})



@login_required
@ad_permission_required("groups", "delete")
def delete_bulk_groups(request):
    results = []

    if request.method == "POST":
        upload = _get_uploaded_file(request, "file", "excel_file")

        if not upload:
            messages.error(request, "Please upload an Excel file.")
        else:
            try:
                results = bulk_delete_groups_from_excel(upload)

                success_count = sum(1 for r in results if r["success"])
                fail_count = len(results) - success_count

                if success_count:
                    messages.success(request, f"{success_count} groups deleted successfully.")
                if fail_count:
                    messages.warning(request, f"{fail_count} rows failed.")
            except Exception as e:
                messages.error(request, f"Bulk delete failed: {e}")

    return render(request, "management/group/delete/bulk_group.html", {"results": results, "mode": "delete"})


RESOURCES = ["users", "computers", "ous", "groups", "reports"]
ACTIONS = ["create", "delete", "modify", "lock", "unlock", "reset", "view_reports"]

ROLE_CHOICES = [
    ("helpdesk", "Help-Desk"),
    ("auditor", "Auditor"),
    ("admin", "Admin"),
    ("custom", "Custom"),
]


def _default_permissions_for_role(role: str) -> dict:
    role = (role or "custom").strip().lower()

    empty = {
        "users": [],
        "computers": [],
        "ous": [],
        "groups": [],
        "reports": [],
    }

    presets = {
        "admin": {
            "users": ["create", "delete", "modify", "lock", "unlock", "reset"],
            "computers": ["create", "delete", "modify", "lock", "unlock", "reset"],
            "ous": ["create", "delete", "modify"],
            "groups": ["create", "delete", "modify"],
            "reports": ["view_reports"],
        },
        "helpdesk": {
            "users": ["create", "modify", "lock", "unlock", "reset"],
            "computers": ["modify", "lock", "unlock"],
            "ous": [],
            "groups": ["modify"],
            "reports": ["view_reports"],
        },
        "auditor": {
            "users": [],
            "computers": [],
            "ous": [],
            "groups": [],
            "reports": ["view_reports"],
        },
        "custom": empty,
    }

    return presets.get(role, empty)


def _permissions_from_request(request, role: str = "custom") -> dict:
    selected = {}

    for resource in RESOURCES:
        values = request.POST.getlist(f"perms_{resource}[]")
        selected[resource] = [value for value in values if value in ACTIONS]

    has_any_permission = any(bool(values) for values in selected.values())

    if not has_any_permission:
        return _default_permissions_for_role(role)

    return selected


def _get_system_users_for_admin():
    User = get_user_model()
    users = User.objects.all().order_by("username")

    items = []

    for user in users:
        profile, _ = HelpdeskProfile.objects.get_or_create(
            user=user,
            defaults={
                "role": "custom",
                "ou": "",
                "scope": "",
                "permissions": _default_permissions_for_role("custom"),
            },
        )

        items.append(
            {
                "user": user,
                "profile": profile,
            }
        )

    return items


def _get_admin_stats(system_users: list[dict]) -> dict:
    total_users = len(system_users)
    helpdesk_users = 0
    admin_users = 0
    auditor_users = 0
    custom_users = 0

    for item in system_users:
        role = (getattr(item["profile"], "role", "") or "").lower()

        if role == "helpdesk":
            helpdesk_users += 1
        elif role == "admin":
            admin_users += 1
        elif role == "auditor":
            auditor_users += 1
        else:
            custom_users += 1

    return {
        "total_users": total_users,
        "helpdesk_users": helpdesk_users,
        "admin_users": admin_users,
        "auditor_users": auditor_users,
        "custom_users": custom_users,
    }


def _get_admin_logs_placeholder():
    """
    Placeholder until you add a real AuditLog model.

    Later, replace this function with:
        return AuditLog.objects.all().order_by("-created_at")[:100]
    """
    return []


@login_required
@system_admin_required
@require_http_methods(["GET"])
def admin_hub(request):
    system_users = _get_system_users_for_admin()
    stats = _get_admin_stats(system_users)

    active_domain = get_active_ldap_settings(request)

    try:
        scope_ous = list_ous(limit=5000)
    except Exception as exc:
        scope_ous = []
        messages.warning(request, f"Could not load OU scope list: {exc}")

    try:
        settings_obj = LdapSettings.get_settings()

        health = {
            "ldap_status": "Not checked",
            "winrm_status": "Not checked",
            "server_uri": settings_obj.server_uri if settings_obj else "",
            "server_name": settings_obj.server_name if settings_obj else "",
            "base_dn": settings_obj.base_dn if settings_obj else "",
            "safe_mode": settings_obj.safe_mode if settings_obj else None,
            "ldap_response_ms": 0,
            "winrm_response_ms": 0,
            "error": "",
        }

    except Exception as exc:
        health = {
            "ldap_status": "Unknown",
            "winrm_status": "Unknown",
            "server_uri": "",
            "server_name": "",
            "base_dn": "",
            "safe_mode": None,
            "ldap_response_ms": 0,
            "winrm_response_ms": 0,
            "error": str(exc),
        }

    context = {
        "resources": RESOURCES,
        "actions": ACTIONS,
        "role_choices": ROLE_CHOICES,
        "system_users": system_users,
        "stats": stats,
        "health": health,
        "scope_ous": scope_ous,
        "ldap_settings_list": LdapSettings.objects.all().order_by("name", "domain_name", "server_uri"),
        "admin_logs": _get_admin_logs_placeholder(),
    }

    return render(request, "admin/admin_hub.html", context)


@login_required
@system_admin_required
@require_POST
def admin_create_helpdesk_user(request):
    username = (request.POST.get("username") or "").strip()
    email = (request.POST.get("email") or "").strip()
    temp_password = request.POST.get("temp_password") or ""
    first_name = (request.POST.get("first_name") or "").strip()
    last_name = (request.POST.get("last_name") or "").strip()
    role = (request.POST.get("role") or "custom").strip().lower()
    ou = (request.POST.get("ou") or "").strip()
    scope = (request.POST.get("scope") or "").strip()
    active_domain = get_active_ldap_settings(request)

    if not username:
        messages.error(request, "Username is required.")
        return redirect("admin_hub")

    if not temp_password:
        messages.error(request, "Temporary password is required.")
        return redirect("admin_hub")

    if role not in dict(ROLE_CHOICES):
        role = "custom"

    User = get_user_model()

    if User.objects.filter(username=username).exists():
        messages.error(request, f"Username already exists: {username}")
        return redirect("admin_hub")

    if role == "custom":
        permissions = _permissions_from_request(request, role)
    else:
        permissions = _default_permissions_for_role(role)

    try:
        user = User.objects.create_user(
            username=username,
            password=temp_password,
        )

        # CustomUser in this project may not have email/first_name/last_name.
        # Set them only when the fields exist.
        if hasattr(user, "email"):
            user.email = email
        if hasattr(user, "first_name"):
            user.first_name = first_name
        if hasattr(user, "last_name"):
            user.last_name = last_name
        user.save()

        domain_permissions = {}
        if active_domain:
            domain_permissions[str(active_domain.id)] = {
                "role": role,
                "scope": scope,
                "permissions": permissions,
            }

        HelpdeskProfile.objects.update_or_create(
            user=user,
            defaults={
                "role": role,
                "ou": ou,
                "scope": scope,
                "permissions": permissions,
                "domain_permissions": domain_permissions,
            },
        )

        messages.success(request, f"System user '{username}' created successfully.")

    except Exception as exc:
        messages.error(request, f"Failed to create system user: {exc}")

    return redirect("admin_hub")


@login_required
@system_admin_required
@require_POST
def admin_assign_roles(request):
    user_id = request.POST.get("user_id")
    role = (request.POST.get("role") or "custom").strip().lower()
    ou = (request.POST.get("ou") or "").strip()
    scope = (request.POST.get("scope") or "").strip()
    active_domain = get_active_ldap_settings(request)

    if role not in dict(ROLE_CHOICES):
        role = "custom"

    if not user_id:
        messages.error(request, "Please select a user.")
        return redirect("admin_hub")

    User = get_user_model()

    try:
        user = User.objects.get(id=user_id)

        profile, _ = HelpdeskProfile.objects.get_or_create(
            user=user,
            defaults={
                "permissions": _default_permissions_for_role(role),
            },
        )

        profile.role = role
        profile.ou = ou
        profile.scope = scope

        if role == "custom":
            if not profile.permissions:
                profile.permissions = _default_permissions_for_role("custom")
        else:
            profile.permissions = _default_permissions_for_role(role)

        domain_permissions = profile.domain_permissions or {}
        if active_domain:
            domain_permissions[str(active_domain.id)] = {
                "role": role,
                "scope": scope,
                "permissions": profile.permissions or _default_permissions_for_role(role),
            }
            profile.domain_permissions = domain_permissions

        profile.save()

        messages.success(request, f"Role, scope, and domain permissions updated for '{user.username}'.")

    except User.DoesNotExist:
        messages.error(request, "Selected user was not found.")

    except Exception as exc:
        messages.error(request, f"Failed to assign role: {exc}")

    return redirect("admin_hub")


@login_required
@system_admin_required
@require_POST
def admin_auth_settings(request):
    """
    Manual Admin Hub health checks.

    test_ldap:
        Runs only the lightweight LDAP test.

    test_winrm:
        Runs only the heavier WinRM / PowerShell test.
    """
    if "test_ldap" in request.POST:
        try:
            result = test_ldap_health()

            status = result.get("status", "Unknown")
            response_ms = result.get("response_ms", 0)

            if status == "Connected":
                messages.success(
                    request,
                    f"LDAP test successful. Response time: {response_ms} ms."
                )
            else:
                messages.error(
                    request,
                    f"LDAP test failed. Status: {status}. {result.get('error', '')}"
                )

        except Exception as exc:
            messages.error(request, f"LDAP test failed: {exc}")

        return redirect("admin_hub")

    if "test_winrm" in request.POST:
        try:
            result = test_winrm_health()

            status = result.get("status", "Unknown")
            response_ms = result.get("response_ms", 0)

            if status == "Connected":
                messages.success(
                    request,
                    f"WinRM test successful. Response time: {response_ms} ms."
                )
            elif status == "Safe Mode":
                messages.warning(
                    request,
                    "WinRM test skipped because Safe Mode is enabled."
                )
            else:
                messages.error(
                    request,
                    f"WinRM test failed. Status: {status}. {result.get('error', '')}"
                )

        except Exception as exc:
            messages.error(request, f"WinRM test failed: {exc}")

        return redirect("admin_hub")

    messages.info(request, "No health check action was selected.")
    return redirect("admin_hub")


@login_required
@system_admin_required
@require_http_methods(["GET"])
def admin_logs(request):
    """
    Placeholder until a persistent AuditLog model is added.
    """
    messages.info(
        request,
        "Admin logs are ready in the UI. Add an AuditLog model later for persistent logging."
    )
    return redirect("admin_hub")
